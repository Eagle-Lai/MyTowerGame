# -*- coding: utf-8 -*-
"""
把 Unity 地图编辑器导出的 .workbuddy/levelmap_export.json 写回 LevelMap.xlsx（M3-7）。

【为什么中间要隔一个文件，而不是让编辑器直接写 xlsx】
  工程里能在编辑器侧读写 Excel 的库（ClosedXML）只存在于 Luban 工具目录；
  把它引进 Assets/Editor 会污染编辑器程序集、也可能和 Unity 自带的
  Newtonsoft/OpenXml 版本打架。所以分工是：
      Unity 窗口只负责"画 + 校验" → 导出 json → 本脚本写 xlsx → Luban 导出成运行时数据。

【谁来调我】
  Unity 窗口「一键写入 xlsx」按钮会自动依次调用：
      本脚本（写 xlsx） → Luban 导出（Assets/ConfigJson） → 回读校验
  所以脚本末尾不再提示"请手工再跑一次 Luban"，编辑器会自己接手。

【自检】写回前再验一遍：行长一致、S/E 各一个、S→E 可达。
  窗口里已经验过，但"导出的文件被手改过"是常有的事，这里再挡一次。
"""
import datetime
import json
import os
import shutil
import sys
from collections import deque

# ---------------------------------------------------------------------------
# ROOT 一律由脚本自身位置推导（<工程根>/.workbuddy/tools/xxx.py 上两级），
# 禁止硬编码绝对路径 —— 否则换盘符/换机器就静默写错地方。
# ---------------------------------------------------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))          # <根>/.workbuddy/tools
ROOT = os.path.dirname(os.path.dirname(_HERE))              # <根>

try:
    import openpyxl
except ImportError:
    raise SystemExit(
        "缺少依赖 openpyxl。请用「装过 openpyxl 的解释器」运行本脚本，或先执行：\n"
        "    \"<python路径>\" -m pip install openpyxl\n"
        "当前解释器：%s" % sys.executable
    )

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

EXPORT = os.path.join(ROOT, ".workbuddy", "levelmap_export.json")
DATAS = os.path.join(ROOT, "Luban", "Config", "Datas")
BACKUP_ROOT = os.path.join(ROOT, ".workbuddy", "backup")

# Luban 源表的约定：第 1 行是字段名，第 2 行是类型，第 3 行是中文注释，数据从第 4 行开始
HEADER_ROW = 1
FIRST_DATA_ROW = 4


def reachable(cells, cols, rows):
    """BFS：障碍 # 与空洞 X 不可通行。返回 (是否可达, 最短步数)。"""
    sr = sc = er = ec = -1
    for r in range(rows):
        for c in range(cols):
            ch = cells[r][c]
            if ch == "S":
                sr, sc = r, c
            elif ch == "E":
                er, ec = r, c
    if sr < 0 or er < 0:
        return False, -1
    dist = {(sr, sc): 0}
    q = deque([(sr, sc)])
    while q:
        r, c = q.popleft()
        if (r, c) == (er, ec):
            return True, dist[(r, c)]
        for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1)):
            nr, nc = r + dr, c + dc
            if 0 <= nr < rows and 0 <= nc < cols and (nr, nc) not in dist and cells[nr][nc] not in "#X":
                dist[(nr, nc)] = dist[(r, c)] + 1
                q.append((nr, nc))
    return False, -1


def header_map(ws):
    """字段名 -> 列号（第 1 行）。"""
    return {ws.cell(row=HEADER_ROW, column=c).value: c for c in range(1, ws.max_column + 1)}


def scene_level_ids():
    """
    读 SceneInfo.xlsx 里已登记的关卡 id 集合。
    【为什么要看】SceneInfo 是"关卡入口表"（决定选关界面能不能进第 N 关），
    LevelMap 只是它的棋盘。只往 LevelMap 加一行、不在 SceneInfo 登记，
    游戏里这张图根本进不去 —— 这种"写了但没人用"的坑必须在写表时就提示出来。
    表不存在或没有 id 列时返回 None（表示无法判断，不报错）。
    """
    path = os.path.join(DATAS, "SceneInfo.xlsx")
    if not os.path.exists(path):
        return None
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return None
    try:
        ws = wb.active
        hdr = header_map(ws)
        if "id" not in hdr:
            return None
        ids = set()
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            v = ws.cell(row=r, column=hdr["id"]).value
            if v is not None:
                ids.add(int(v))
        return ids
    finally:
        wb.close()


def main():
    if not os.path.exists(EXPORT):
        raise SystemExit("找不到导出文件：%s\n请先在 Unity 里用「Tools ▸ 塔防 ▸ 地图编辑器」导出。" % EXPORT)

    with open(EXPORT, "r", encoding="utf-8") as f:
        data = json.load(f)

    level_id = int(data["levelId"])
    cols = int(data["cols"])
    rows = int(data["rows"])
    cells = list(data.get("cells") or [])
    desc = data.get("desc") or ""

    print("导出内容：关卡 %d，%d×%d" % (level_id, cols, rows))
    for line in cells:
        print("   " + line)

    # ---------- 自检 ----------
    problems = []
    if len(cells) != rows:
        problems.append("行数 %d != rows %d" % (len(cells), rows))
    for i, line in enumerate(cells):
        if len(line) != cols:
            problems.append("第 %d 行长度 %d != cols %d" % (i + 1, len(line), cols))
    s_count = sum(line.count("S") for line in cells)
    e_count = sum(line.count("E") for line in cells)
    if s_count != 1:
        problems.append("起点 S 有 %d 个（必须 1 个）" % s_count)
    if e_count != 1:
        problems.append("终点 E 有 %d 个（必须 1 个）" % e_count)
    if not problems:
        ok, plen = reachable(cells, cols, rows)
        if not ok:
            problems.append("S→E 不可达")
        else:
            buildable = sum(line.count(".") for line in cells)
            print("\n校验通过：可建造格 %d，最短路径 %d 格" % (buildable, plen))

    if problems:
        for p in problems:
            print("  x " + p)
        raise SystemExit("校验未通过，未写文件")

    # ---------- 写回 ----------
    xlsx_path = os.path.join(DATAS, "LevelMap.xlsx")
    if not os.path.exists(xlsx_path):
        raise SystemExit("找不到源表：%s" % xlsx_path)

    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = os.path.join(BACKUP_ROOT, stamp, "Config", "Datas")
    os.makedirs(dst, exist_ok=True)
    shutil.copy2(xlsx_path, os.path.join(dst, "LevelMap.xlsx"))
    print("已备份到: %s" % dst)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    hdr = header_map(ws)
    for need in ("id", "levelId", "cols", "rows", "cells"):
        if need not in hdr:
            raise SystemExit("LevelMap.xlsx 缺少字段列 '%s'（表头是否被改过？）" % need)

    target = None
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(row=r, column=hdr["id"]).value
        if v is not None and int(v) == level_id:
            target = r
            break

    created = False
    if target is None:
        target = ws.max_row + 1
        created = True
        ws.cell(row=target, column=hdr["id"]).value = level_id
        ws.cell(row=target, column=hdr["levelId"]).value = level_id

    ws.cell(row=target, column=hdr["cols"]).value = cols
    ws.cell(row=target, column=hdr["rows"]).value = rows
    ws.cell(row=target, column=hdr["cells"]).value = "|".join(cells)
    if "desc" in hdr:
        # 无条件写入：清空备注也要能同步，否则旧备注会一直赖在表里
        ws.cell(row=target, column=hdr["desc"]).value = desc

    wb.save(xlsx_path)
    print("已%s LevelMap.xlsx 第 %d 行（关卡 %d，%d×%d）"
          % ("新增" if created else "更新", target, level_id, cols, rows))

    # ---------- 软提示：SceneInfo 是否登记过这一关 ----------
    ids = scene_level_ids()
    if ids is not None and level_id not in ids:
        print("")
        print("[提示] SceneInfo.xlsx 里没有 id=%d 这一关 —— 棋盘写好了，但游戏里还进不去。" % level_id)
        print("       需要同时在 Luban/Config/Datas/SceneInfo.xlsx 登记该关（关卡入口/波次/经济等）。")

    print("")
    print("写表完成。Unity 会接着自动跑 Luban 导出并在编辑器里回读校验。")


if __name__ == "__main__":
    main()

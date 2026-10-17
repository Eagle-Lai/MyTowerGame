# -*- coding: utf-8 -*-
"""
新增关卡 16~20 的完整配置（EnemyList / RoundData / SceneInfo），配合已画好的 LevelMap 16~20。

【与 add_levels_9_15_20261017.py 的关系】同一套口径，只是接着往后延：
  · 波次难度仍锚定第 40 波（59681 有效血量）× 1.008^(w-40)，所以 75 波处正好接上
    上次的 78878，100 波约 96260 —— 曲线是连续的，不是重新起一条。
  · 初始金币继续 +150/关（2600 起、L15=3500 → L16=3650 … L20=4250），与敌人强度同步。
  · 每关 5 回合：L16=76~80, L17=81~85, L18=86~90, L19=91~95, L20=96~100。

【只追加/覆盖自己这段】脚本可重复运行（按 id 找行，找到就覆盖，不会重复追加）。
【地图事实】16~20 都是 16/17×9 的小障碍图，备注按实际棋盘写。
"""
import datetime
import os
import shutil
import sys

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

_HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(_HERE))
DATAS = os.path.join(ROOT, "Luban", "Config", "Datas")
BACKUP_ROOT = os.path.join(ROOT, ".workbuddy", "backup")

FIRST_DATA_ROW = 4

FIRST_WAVE = 76
LAST_WAVE = 100
WAVE_EHP_AT_40 = 59681.0     # 与 9~15 那批完全同一锚点，保证曲线连续
WAVE_GROWTH = 1.008
MAX_COUNT = 46
BOSSES = [129, 130, 131]
ELITE = 132

# 关卡号 -> (初始金币, 初始生命, 难度, 备注)
NEW_LEVELS = {
    16: (3650, 12, 3, "M16 新图：16×9，col8 竖墙 3 格把左右隔开，S 左中 → E 右中"),
    17: (3800, 12, 3, "M17 新图：17×9，中央 3 格横墙，S 上中 → E 下中"),
    18: (3950, 12, 3, "M18 新图：16×9，col8 竖墙 3 格，S 右中 → E 左中（16 的镜像）"),
    19: (4100, 12, 3, "M19 新图：16×9，左下角 3 格横墙，S 左上 → E 左下"),
    20: (4250, 12, 3, "M20 新图：16×9，S/E 两端各贴一格障碍，中间全开阔"),
}


def header_map(ws):
    return {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}


def last_data_row(ws, id_col):
    last = FIRST_DATA_ROW - 1
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(row=r, column=id_col).value is not None:
            last = r
    return last


def find_or_append_row(ws, id_col, target_id):
    """按 id 找行；没有就追加。保证脚本可重复运行。"""
    last = FIRST_DATA_ROW - 1
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(row=r, column=id_col).value
        if v is None:
            continue
        last = r
        if int(v) == target_id:
            return r
    r = last + 1
    ws.cell(row=r, column=id_col).value = target_id
    return r


def load_enemies():
    wb = openpyxl.load_workbook(os.path.join(DATAS, "EnemyData.xlsx"), data_only=True)
    ws = wb.active
    hdr = header_map(ws)
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        eid = ws.cell(row=r, column=hdr["id"]).value
        if eid is None:
            continue
        eid = int(eid)
        hp = float(ws.cell(row=r, column=hdr["hp"]).value or 0)
        ar = float(ws.cell(row=r, column=hdr["armor"]).value or 0)
        rows.append((eid, hp / max(0.05, 1.0 - ar)))
    wb.close()
    return rows


def tiers_of(rows):
    normal = [x for x in rows if x[0] not in BOSSES and x[0] != ELITE]
    normal.sort(key=lambda x: x[1])
    n = len(normal)
    return [[e for e, _ in normal[i * n // 4:(i + 1) * n // 4]] for i in range(4)]


def build_wave(w, tiers, eff):
    target = WAVE_EHP_AT_40 * (WAVE_GROWTH ** (w - 40))
    pool = tiers[3]

    ids, acc, i = [], 0.0, 0
    while acc < target and len(ids) < MAX_COUNT:
        eid = pool[(i * 7 + w * 3) % len(pool)]
        ids.append(eid)
        acc += eff[eid]
        i += 1

    if w % 2 == 0:
        ids.append(ELITE)
        acc += eff[ELITE]
    if w % 5 == 0:
        boss = BOSSES[(w // 5) % len(BOSSES)]
        ids.append(boss)
        acc += eff[boss]

    enemy_interval = max(170, 560 - 12 * ((w - 1) // 5) * 4 - 8 * ((w - 1) % 5))
    desc = "第 %d 波：%d 只 / 目标血量 %.0f（单怪间隔 %.2fs）" % (w, len(ids), target, enemy_interval / 1000.0)
    return ids, enemy_interval, desc


def main():
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = os.path.join(BACKUP_ROOT, stamp, "Config", "Datas")
    os.makedirs(dst, exist_ok=True)
    for n in ("EnemyList.xlsx", "RoundData.xlsx", "SceneInfo.xlsx"):
        shutil.copy2(os.path.join(DATAS, n), os.path.join(dst, n))
    print("已备份到:", dst)

    enemies = load_enemies()
    tiers = tiers_of(enemies)
    eff = {e: v for e, v in enemies}

    # ---------- EnemyList：76~100 ----------
    path = os.path.join(DATAS, "EnemyList.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = header_map(ws)
    for w in range(FIRST_WAVE, LAST_WAVE + 1):
        ids, ei, desc = build_wave(w, tiers, eff)
        row = find_or_append_row(ws, hdr["id"], w)
        ws.cell(row=row, column=hdr["id"]).value = w
        ws.cell(row=row, column=hdr["EnemyIndexs"]).value = ",".join(str(x) for x in ids)
        ws.cell(row=row, column=hdr["interval"]).value = 3000
        ws.cell(row=row, column=hdr["enemyInterval"]).value = ei
        ws.cell(row=row, column=hdr["desc"]).value = desc
    wb.save(path)
    print("EnemyList：已写入 %d~%d 波" % (FIRST_WAVE, LAST_WAVE))

    # ---------- RoundData：76~100 ----------
    path = os.path.join(DATAS, "RoundData.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = header_map(ws)
    for rnd in range(FIRST_WAVE, LAST_WAVE + 1):
        row = find_or_append_row(ws, hdr["id"], rnd)
        ws.cell(row=row, column=hdr["id"]).value = rnd
        ws.cell(row=row, column=hdr["EnemyIndexs"]).value = str(rnd)
        ws.cell(row=row, column=hdr["interval"]).value = 1500
        ws.cell(row=row, column=hdr["name"]).value = "第 %d 回合" % rnd
        ws.cell(row=row, column=hdr["rewardGold"]).value = 30 + 4 * (rnd - 1)
    wb.save(path)
    print("RoundData：已写入 %d~%d 回合" % (FIRST_WAVE, LAST_WAVE))

    # ---------- SceneInfo：16~20 ----------
    path = os.path.join(DATAS, "SceneInfo.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = header_map(ws)
    id_c = hdr["id"]

    template = None
    for rr in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(row=rr, column=id_c).value
        if v is not None and int(v) == 8:
            template = rr
            break
    if template is None:
        raise SystemExit("SceneInfo.xlsx 里找不到第 8 关，无法取模板行")

    def cell(rid):
        return ws.cell(row=template, column=hdr[rid]).value

    for lid in sorted(NEW_LEVELS):
        gold, hp, diff, desc = NEW_LEVELS[lid]
        rounds = list(range((lid - 9) * 5 + 41, (lid - 9) * 5 + 46))
        row = find_or_append_row(ws, id_c, lid)
        ws.cell(row=row, column=hdr["id"]).value = lid
        ws.cell(row=row, column=hdr["RoundList"]).value = ",".join(str(x) for x in rounds)
        ws.cell(row=row, column=hdr["CameraPosition"]).value = cell("CameraPosition")
        ws.cell(row=row, column=hdr["CameraRotration"]).value = cell("CameraRotration")
        ws.cell(row=row, column=hdr["MapName"]).value = cell("MapName")
        ws.cell(row=row, column=hdr["EenemyPosition"]).value = cell("EenemyPosition")
        ws.cell(row=row, column=hdr["MapPosition"]).value = cell("MapPosition")
        ws.cell(row=row, column=hdr["PointScale"]).value = cell("PointScale")
        ws.cell(row=row, column=hdr["name"]).value = "关卡 %d" % lid
        ws.cell(row=row, column=hdr["mapId"]).value = lid
        ws.cell(row=row, column=hdr["initialGold"]).value = gold
        ws.cell(row=row, column=hdr["initialHp"]).value = hp
        ws.cell(row=row, column=hdr["waveIntervalMs"]).value = 3000
        ws.cell(row=row, column=hdr["nightSightMs"]).value = 0
        ws.cell(row=row, column=hdr["difficulty"]).value = diff
        ws.cell(row=row, column=hdr["desc"]).value = desc
    wb.save(path)
    print("SceneInfo：已写入 16~20 关（每关 5 回合）")

    # ---------- 自检 ----------
    wb = openpyxl.load_workbook(os.path.join(DATAS, "LevelMap.xlsx"), data_only=True)
    ws = wb.active
    h = header_map(ws)
    map_ids = set()
    for rr in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(row=rr, column=h["id"]).value
        if v is not None:
            map_ids.add(int(v))
    wb.close()

    print("\n【自检】")
    print("  关卡  mapId对齐  回合区间     首/末波目标血量   金币  生命")
    for lid in sorted(NEW_LEVELS):
        gold, hp, diff, _ = NEW_LEVELS[lid]
        rounds = list(range((lid - 9) * 5 + 41, (lid - 9) * 5 + 46))
        first = WAVE_EHP_AT_40 * (WAVE_GROWTH ** (rounds[0] - 40))
        last = WAVE_EHP_AT_40 * (WAVE_GROWTH ** (rounds[-1] - 40))
        ok = "OK" if lid in map_ids else "缺失!"
        print("   %2d    %-6s   %3d~%3d   %8.0f/%-8.0f  %4d  %2d"
              % (lid, ok, rounds[0], rounds[-1], first, last, gold, hp))


if __name__ == "__main__":
    main()

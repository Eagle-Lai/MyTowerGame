#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""一次性手工改表（L01 波次奖励，2026-11-01 用户口径，可传增量）。

★ 口径：**替换式**。最终值 = **改动前原值** + flat；W5/W10/W15/W20 = 原值 + (flat + bonus)。
  基线固定取 `RoundData.xlsx.bak_20261101_goldhp`（2026-11-01 之前的原值），
  所以同一条命令重复执行结果一致，不会越跑越高；调增量只需改参数，不用改代码。

只改 `Levels/L01/RoundData.xlsx` 的 `rewardGold` 一列 25 行，血量与其他表一律不动。

★ 波次顺序以 `SceneInfo.RoundList` 为准（1..20, 401-405 ⇒ W21-W25 = roundId 401-405）；
   `RoundData.xlsx` 的物理行顺序不是波次顺序。

用法：
  python _l01_wave_gold_revise_20261101.py values [flat] [bonus]
  python _l01_wave_gold_revise_20261101.py apply  [flat] [bonus]
  python _l01_wave_gold_revise_20261101.py verify [flat] [bonus]
  # 缺省 flat=7 bonus=20（首次修订的数值）；每波+3、里程碑再+5 则传 3 5
"""
import pathlib
import shutil
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
BASELINE = L01 / 'RoundData.xlsx.bak_20261101_goldhp'   # 2026-11-01 改动之前的原值（固定基线）
CUR = L01 / 'RoundData.xlsx'
GOLD_BONUS_WAVES = {5, 10, 15, 20}


def heads(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def wave_order():
    wb = openpyxl.load_workbook(L01 / 'SceneInfo.xlsx', data_only=True)
    ws = wb['Sheet1']
    h = heads(ws)
    order = [int(x) for x in str(ws.cell(4, h['RoundList']).value).split(',') if str(x).strip()]
    wb.close()
    return order


def read_gold(path):
    """{roundId: (row, rewardGold)}；.bak 先复制成 .xlsx 再读（openpyxl 只认扩展名）。"""
    tmp = ROOT / '.workbuddy/tools/__pycache__/bakcheck'
    tmp.mkdir(parents=True, exist_ok=True)
    src = path
    if path.suffix != '.xlsx':
        src = tmp / f'RoundData_{path.name[-12:]}.xlsx'
        shutil.copy2(path, src)
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[wb.sheetnames[0]]
    h = heads(ws)
    out = {}
    for r in range(4, ws.max_row + 1):
        rid = ws.cell(r, h['id']).value
        if rid is None:
            continue
        out[int(rid)] = (r, int(ws.cell(r, h['rewardGold']).value))
    wb.close()
    return out


def main():
    args = [a for a in sys.argv[1:]]
    mode = args[0] if args else 'values'
    flat = int(args[1]) if len(args) > 1 else 7
    bonus = int(args[2]) if len(args) > 2 else 20

    order = wave_order()
    assert BASELINE.exists(), f'缺少原值基线 {BASELINE.name}'
    base = read_gold(BASELINE)
    cur = read_gold(CUR)
    plan = {}
    for w, rid in enumerate(order, 1):
        add = flat + (bonus if w in GOLD_BONUS_WAVES else 0)
        plan[rid] = (base[rid][1], base[rid][1] + add)

    if mode in ('values', 'verify'):
        bad = 0
        print(f'口径：原值 +{flat}（W5/10/15/20 再 +{bonus} = 原值 +{flat + bonus}）')
        print('波 | roundId | 原值 | 当前 | 目标 | 相对原值')
        for w, rid in enumerate(order, 1):
            o, n = plan[rid]
            print(f'W{w:<2} | {rid:<7} | {o:>4} | {cur[rid][1]:>4} | {n:>4} | +{n - o}')
            if mode == 'verify' and cur[rid][1] != n:
                bad += 1
        so = sum(plan[r][0] for r in order)
        sn = sum(plan[r][1] for r in order)
        print(f'合计 原值 {so} -> 目标 {sn}  (Δ+{sn - so})')
        if mode == 'verify':
            if bad:
                print(f'❌ 校验失败：{bad} 行与目标不符')
                return 1
            print('✅ 校验通过：25 行 rewardGold 均 = 原值 + 增量')
        return 0

    assert mode == 'apply', f'未知模式 {mode}'

    if all(cur[rid][1] == plan[rid][1] for rid in order):
        print('⏭ 已是目标值，未重复写入')
        return 0

    # 备份"本次改动前"的版本，文件名带上当时的每波增量（相对原值），便于回滚链辨认
    cur_flat = cur[order[0]][1] - base[order[0]][1]
    bak = L01 / f'RoundData.xlsx.bak_20261101_gold_prev{cur_flat}'
    if not bak.exists():
        shutil.copy2(CUR, bak)
        print(f'备份（改动前，每波 +{cur_flat} 版本）-> {bak.name}')

    wb = openpyxl.load_workbook(CUR)
    ws = wb[wb.sheetnames[0]]
    h = heads(ws)
    for rid in order:
        ws.cell(cur[rid][0], h['rewardGold']).value = plan[rid][1]
    wb.save(CUR)

    print(f'已写入 25 行 rewardGold（原值 +{flat}，W5/10/15/20 原值 +{flat + bonus}）：')
    for w, rid in enumerate(order, 1):
        o, n = plan[rid]
        print(f'  W{w:<2} id={rid:<4} {o:>3} -> {n}   (+{n - o})')
    print(f"合计 {sum(plan[r][0] for r in order)} -> {sum(plan[r][1] for r in order)}")
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

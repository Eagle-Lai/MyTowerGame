"""向 TowerInfo.xlsx 追加 type=2 / type=3 塔的配置行（A-1）。

表格结构（保持不变）：
  第 1 行 ##var  （字段名）
  第 2 行 ##type （类型）
  第 3 行 中文说明
  第 4 行起 数据

现有：id 1/2/3 = type=1(Normal) lv1/lv2/lv3
追加：id 4/5/6 = type=2(Power) lv1/2/3，id 7/8/9 = type=3(Retard) lv1/2/3

列序（A..S）：
  A(##) B(id) C(type) D(name) E(resName) F(level) G(radius) H(power) I(CD)
  J(prices) K(bulletId) L(targetMode) M(searchIntervalMs) N(rotateSpeed)
  O(upgradeTo) P(sellPrice) Q(effectType) R(effectValue) S(desc)
"""
import shutil, os
import openpyxl

SRC = 'Luban/Config/Datas/TowerInfo.xlsx'
BAK = 'Luban/Config/Datas/_backup_TowerInfo_before_M2.xlsx'

# 每行按列序 A..S 排（A 留空）
ROWS = [
    # ---- type=2 强力塔 ----
    [None, 4, 2, 'PowerTower', 'Tower_Power', 1, 3, 25, 600, 35, 1, 0, 100, 360, 5, 24, 2, 0, None],
    [None, 5, 2, 'PowerTower', 'Tower_Power', 2, 4, 32, 500, 55, 1, 0, 100, 360, 6, 62, 2, 0, None],
    [None, 6, 2, 'PowerTower', 'Tower_Power', 3, 5, 45, 400, 85, 1, 0, 100, 360, 0, 122, 2, 0, None],
    # ---- type=3 减速塔 ----
    [None, 7, 3, 'RetardTower', 'Tower_Retard', 1, 4, 5, 400, 30, 1, 0, 100, 360, 8, 21, 1, 0.5, None],
    [None, 8, 3, 'RetardTower', 'Tower_Retard', 2, 5, 7, 350, 45, 1, 0, 100, 360, 9, 52, 1, 0.6, None],
    [None, 9, 3, 'RetardTower', 'Tower_Retard', 3, 6, 10, 300, 70, 1, 0, 100, 360, 0, 102, 1, 0.7, None],
]


def main():
    if not os.path.exists(BAK):
        shutil.copy2(SRC, BAK)
        print('已备份 →', BAK)

    wb = openpyxl.load_workbook(SRC)
    ws = wb['Sheet1']

    # 找到现有数据的最后一行（跳过空行）
    last = 3
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=2).value not in (None, ''):
            last = r

    start = last + 1
    for i, rowvals in enumerate(ROWS):
        r = start + i
        for c, val in enumerate(rowvals, start=1):
            if val is None:
                continue
            ws.cell(row=r, column=c).value = val

    wb.save(SRC)
    print(f'已写入 {len(ROWS)} 行 → 第 {start}~{start + len(ROWS) - 1} 行')

    # 回读校验
    wb2 = openpyxl.load_workbook(SRC)
    ws2 = wb2['Sheet1']
    for r in range(4, start + len(ROWS)):
        vals = [ws2.cell(row=r, column=c).value for c in range(2, 19)]
        print(r, vals)


main()

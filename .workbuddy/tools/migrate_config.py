# -*- coding: utf-8 -*-
"""
Luban 配置表迁移脚本 —— FreeTower M0

作用：
  1. 备份 Luban/Config 到 .workbuddy/backup/<时间戳>/
  2. 扩展现有 4 张表（EnemyData / TowerInfo / RoundData / SceneInfo）的列
  3. 修正 TowerInfo 的两处数据问题（type 随 level 递增、path 语义）
  4. 新建 3 张表（BulletData / LevelMap / Global）
  5. 在 __tables__.xlsx 中登记新表

Luban 表格式约定（实测确认）：
  A 列 = 标记列（"##" / "##type" / "##"）
  第 1 行 = 字段名，第 2 行 = 类型，第 3 行 = 中文注释，第 4 行起 = 数据
"""
import os
import shutil
import datetime

import openpyxl

ROOT = r'D:\FreedomTower'
DATAS = os.path.join(ROOT, 'Luban', 'Config', 'Datas')
BACKUP_ROOT = os.path.join(ROOT, '.workbuddy', 'backup')


# ----------------------------------------------------------------------
# 工具
# ----------------------------------------------------------------------

def last_header_col(ws):
    """返回第 1 行最后一个非空列号"""
    last = 0
    for c in range(1, ws.max_column + 2):
        v = ws.cell(1, c).value
        if v is not None and str(v).strip() != '':
            last = c
    return last


def header_names(ws):
    return [ws.cell(1, c).value for c in range(1, last_header_col(ws) + 1)]


def is_blank(v):
    """None 或纯空白都算空（避免把空字符串当有效数据）"""
    return v is None or str(v).strip() == ''


def append_columns(ws, cols):
    """在表尾追加列。cols = [(字段名, 类型, 注释, 逐行取值函数或常量), ...]"""
    start = last_header_col(ws) + 1
    # 只认「前置列里有真实数据」的行作为数据行，避免把尾部空白行当数据
    n_data = 3
    for r in range(4, ws.max_row + 1):
        if any(not is_blank(ws.cell(r, c).value) for c in range(2, start)):
            n_data = r
    for i, (name, typ, comment, value) in enumerate(cols):
        c = start + i
        ws.cell(1, c).value = name
        ws.cell(2, c).value = typ
        ws.cell(3, c).value = comment
        for r in range(4, n_data + 1):
            v = value(r) if callable(value) else value
            if not is_blank(v):
                ws.cell(r, c).value = v
    print('      数据行数: %d（第 4 ~ %d 行），新增列从第 %d 列开始' % (n_data - 3, n_data, start))
    return start


def find_col(ws, name):
    for c in range(1, last_header_col(ws) + 1):
        if ws.cell(1, c).value == name:
            return c
    raise KeyError('找不到字段: ' + name)


def new_table(path, cols, rows):
    """新建一张表。cols = [(名, 类型, 注释)], rows = [[值...], ...]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.cell(1, 1).value = '##'
    ws.cell(2, 1).value = '##type'
    ws.cell(3, 1).value = '##'
    for i, (name, typ, comment) in enumerate(cols):
        c = 2 + i
        ws.cell(1, c).value = name
        ws.cell(2, c).value = typ
        ws.cell(3, c).value = comment
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            if v is not None and v != '':
                ws.cell(4 + ri, 2 + ci).value = v
    # 预留空 sheet 无关紧要，但要保证只有 Sheet1 被 Luban 读
    wb.save(path)
    print('  [新建] %s  (%d 列 x %d 行数据)' % (os.path.basename(path), len(cols), len(rows)))


def load(name):
    p = os.path.join(DATAS, name)
    wb = openpyxl.load_workbook(p)
    ws = wb['Sheet1']
    # 数据区右侧存在手工编辑残留的合并单元格（如 L7:M7），
    # 会挡住追加列（MergedCell 只读）。这些区域无实际数据，直接解除合并。
    stray = [str(r) for r in ws.merged_cells.ranges if r.min_row >= 4]
    for rng in stray:
        ws.unmerge_cells(rng)
    if stray:
        print('      解除数据区多余合并: %s' % stray)
    return wb, ws, p


def save(wb, p, tag):
    wb.save(p)
    print('  [扩展] %s  → %s' % (os.path.basename(p), tag))


# ----------------------------------------------------------------------
# 0. 备份
# ----------------------------------------------------------------------

ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
dst = os.path.join(BACKUP_ROOT, ts, 'Config')
os.makedirs(os.path.dirname(dst), exist_ok=True)
shutil.copytree(os.path.join(ROOT, 'Luban', 'Config'), dst)
print('[备份] Luban/Config  →  %s' % dst)
print()


# ----------------------------------------------------------------------
# 1. EnemyData.xlsx  (TBEnemyData)
# ----------------------------------------------------------------------

print('[1/7] EnemyData.xlsx')
wb, ws, p = load('EnemyData.xlsx')

# 现有：id, speed, hp, name, type, interval
# speed 原为 3D 世界单位/秒（12/24/6...），2D 下 CellSize=1 格，
# 按 ÷8 换算为「格/秒」，否则敌人会以 12 格/秒横穿棋盘。
SPEED_DIV = 8.0
speed_col = find_col(ws, 'speed')
for r in range(4, ws.max_row + 1):
    v = ws.cell(r, speed_col).value
    if isinstance(v, (int, float)) and v:
        ws.cell(r, speed_col).value = round(v / SPEED_DIV, 3)

type_col = find_col(ws, 'type')


def by_type(mapping, default):
    def f(r):
        t = ws.cell(r, type_col).value
        return mapping.get(t, default(r) if callable(default) else default)
    return f


append_columns(ws, [
    ('resName', 'string', '战斗预制体资源名（对应 ResTable 逻辑名）', 'Enemy_Rat'),
    ('armor', 'float', '护甲减伤 0~1（最终伤害=伤害*(1-armor)）',
     by_type({1: 0.0, 2: 0.0, 3: 0.2}, 0.0)),
    ('reward', 'int', '击杀奖励金币', by_type({1: 5, 2: 8, 3: 15}, 5)),
    ('damageToPlayer', 'int', '漏怪扣除玩家生命', 1),
    ('scale', 'float', '显示缩放', 1.0),
    ('bodyRadius', 'float', '受击半径（单位：格，索敌与命中判定用）', 0.4),
    ('isFlying', 'int', '是否飞行 0否/1是（M0 全 0；M1 可改 bool）', 0),
    ('animController', 'string', 'AnimatorController 家族名（对应 _Common/Animations/<名>/Controller）', 'Rat'),
    ('desc', 'string', '备注', ''),
])
save(wb, p, 'speed÷8 换算为格/秒，新增 9 列')


# ----------------------------------------------------------------------
# 2. TowerInfo.xlsx  (TBTowerInfo)
# ----------------------------------------------------------------------

print('[2/7] TowerInfo.xlsx')
wb, ws, p = load('TowerInfo.xlsx')

# 修正 1：type 原来的值是 1/2/3 且随 level 递增 —— 语义错误。
#         三行都是普通塔，type 应统一为 1（塔类型），level 才是 1/2/3。
c_type = find_col(ws, 'type')
ws.cell(2, c_type).value = 'int'          # 类型声明：string → int
for r in range(4, ws.max_row + 1):
    if ws.cell(r, c_type).value is not None:
        ws.cell(r, c_type).value = 1

# 修正 2：path 改名为 resName，语义从「Resources 路径」改为「AB 逻辑名」
c_path = find_col(ws, 'path')
ws.cell(1, c_path).value = 'resName'
ws.cell(2, c_path).value = 'string'
ws.cell(3, c_path).value = '资源名（对应 ResTable 逻辑名）'
for r in range(4, ws.max_row + 1):
    if ws.cell(r, c_path).value is not None:
        ws.cell(r, c_path).value = 'Tower_Normal'

# 修正 3：prices 注释明确为「增量价」，并补 sellPrice
c_price = find_col(ws, 'prices')
level_col = find_col(ws, 'level')

_cumulative = {}


def sell_price(r):
    """按累计成本 70% 返还"""
    lv = ws.cell(r, level_col).value
    if lv is None:
        return None
    total = 0
    for rr in range(4, ws.max_row + 1):
        l2 = ws.cell(rr, level_col).value
        if isinstance(l2, int) and l2 <= lv:
            v = ws.cell(rr, c_price).value
            if isinstance(v, (int, float)):
                total += v
    return int(total * 0.7)


append_columns(ws, [
    ('bulletId', 'int', '发射的子弹 id → TBBulletData', 1),
    ('targetMode', 'int', '索敌策略 0路径最靠前/1最近/2最高血', 0),
    ('searchIntervalMs', 'int', '索敌间隔（毫秒，节流用）', 100),
    ('rotateSpeed', 'float', '炮管转向速度（度/秒）', 360.0),
    ('upgradeTo', 'int', '下一等级 id，0=满级',
     lambda r: 2 if ws.cell(r, level_col).value == 1 else (3 if ws.cell(r, level_col).value == 2 else 0)),
    ('sellPrice', 'int', '出售返还金币（累计成本 70%）', sell_price),
    ('effectType', 'int', '特殊效果类型 0无/1减速/2持续伤害', 0),
    ('effectValue', 'float', '特殊效果参数', 0.0),
    ('desc', 'string', '备注', ''),
])
save(wb, p, 'type 改 int 并统一为 1、path→resName、新增 9 列')


# ----------------------------------------------------------------------
# 3. RoundData.xlsx  (TBRoundData)
# ----------------------------------------------------------------------

print('[3/7] RoundData.xlsx')
wb, ws, p = load('RoundData.xlsx')
append_columns(ws, [
    ('name', 'string', '回合名（展示用）', lambda r: '第 %s 回合' % ws.cell(r, 2).value),
    ('rewardGold', 'int', '本回合通关奖励金币', 20),
])
save(wb, p, '新增 2 列')


# ----------------------------------------------------------------------
# 4. SceneInfo.xlsx  (TBSceneInfo) —— 语义上即「关卡表」
# ----------------------------------------------------------------------

print('[4/7] SceneInfo.xlsx')
wb, ws, p = load('SceneInfo.xlsx')
append_columns(ws, [
    ('name', 'string', '关卡名', lambda r: '关卡 %s' % ws.cell(r, 2).value),
    ('mapId', 'int', '棋盘布局 id → TBLevelMap', lambda r: ws.cell(r, 2).value),
    ('initialGold', 'int', '初始金币', 100),
    ('initialHp', 'int', '初始生命', 10),
    ('waveIntervalMs', 'int', '回合间准备时间（毫秒）', 3000),
    ('nightSightMs', 'int', '（预留）视野遮蔽，M0 未用', 0),
    ('difficulty', 'int', '难度标签 1~3', lambda r: min(3, max(1, (ws.cell(r, 2).value or 1)))),
    ('desc', 'string', '备注（CameraPosition 等 3D 字段已废弃，代码不再读取）', ''),
])
save(wb, p, '新增 8 列（3D 相机/坐标字段保留但不再使用）')


# ----------------------------------------------------------------------
# 5. 新建 BulletData.xlsx  (TBBulletData)
# ----------------------------------------------------------------------

print('[5/7] BulletData.xlsx（新建）')
new_table(os.path.join(DATAS, 'BulletData.xlsx'), [
    ('id', 'int', '主键'),
    ('name', 'string', '展示名'),
    ('resName', 'string', '资源名（对应 ResTable 逻辑名）'),
    ('speed', 'float', '飞行速度（格/秒）'),
    ('hitRadius', 'float', '命中判定半径（格）'),
    ('lifeTimeMs', 'int', '最大存活时间（毫秒），超时回收'),
    ('pierce', 'int', '穿透数量 0=命中即消失'),
    ('aoeRadius', 'float', '爆炸半径（格）0=单体'),
    ('effectType', 'int', '效果类型 0无/1减速/2持续伤害'),
    ('effectValue', 'float', '效果参数'),
    ('scale', 'float', '显示缩放'),
    ('desc', 'string', '备注'),
], [
    [1, '普通子弹', 'Bullet_Normal', 12.0, 0.25, 3000, 0, 0.0, 0, 0.0, 1.0, 'M0 唯一子弹，直线追踪'],
])


# ----------------------------------------------------------------------
# 6. 新建 LevelMap.xlsx  (TBLevelMap)
# ----------------------------------------------------------------------

print('[6/7] LevelMap.xlsx（新建）')

# 16 列 x 9 行棋盘。CellSize = 1 格，正交相机 size=5 → 可视 17.8 x 10 格，留有余量。
# 字符集：S 起点 / E 终点 / . 可建造地砖 / # 障碍 / X 空洞 / P 装饰地砖
MAP1 = [
    'S...............',
    '................',
    '...###....###...',
    '...###....###...',
    '................',
    '................',
    '...###....###...',
    '...###....###...',
    '...............E',
]
assert all(len(r) == 16 for r in MAP1), '地图行长必须为 16'
assert len(MAP1) == 9, '地图必须为 9 行'

new_table(os.path.join(DATAS, 'LevelMap.xlsx'), [
    ('id', 'int', '主键'),
    ('levelId', 'int', '所属关卡 id → TBSceneInfo'),
    ('cols', 'int', '列数'),
    ('rows', 'int', '行数'),
    ('cells', '(list#sep=|),string', '每行一个字符串（用 | 分隔），长度为 cols'),
    ('desc', 'string', '备注'),
], [
    [1, 1, 16, 9, '|'.join(MAP1), 'M0 测试棋盘：开阔场地 + 4 组方形障碍，保证多路径可绕'],
])


# ----------------------------------------------------------------------
# 7. 新建 Global.xlsx  (TBGlobal)
# ----------------------------------------------------------------------

print('[7/7] Global.xlsx（新建）')
new_table(os.path.join(DATAS, 'Global.xlsx'), [
    ('id', 'int', '主键（单行表，固定 1）'),
    ('cellSize', 'float', '棋盘格子边长（世界单位）'),
    ('boardOriginPadding', 'float', '棋盘相对相机中心的留白（世界单位）'),
    ('enemyPoolSize', 'int', '怪物对象池初始容量'),
    ('bulletPoolSize', 'int', '子弹对象池初始容量'),
    ('maxEnemyAlive', 'int', '同屏怪物上限（性能保护）'),
    ('defaultReward', 'int', '未配 reward 时的默认击杀奖励'),
    ('defaultDamageToPlayer', 'int', '未配 damageToPlayer 时的默认漏怪伤害'),
    ('defaultBodyRadius', 'float', '未配 bodyRadius 时的默认受击半径'),
    ('defaultSearchIntervalMs', 'int', '未配 searchIntervalMs 时的默认索敌间隔'),
    ('targetFrameRate', 'int', '目标帧率'),
    ('deathRecycleDelayMs', 'int', '死亡动画播完后延迟回收的时间'),
    ('checkOfflinePerFrame', 'int', '每帧最多重算路径次数（防止卡顿）'),
    ('showDebugLog', 'int', '是否打印战斗自检日志 0否/1是'),
], [
    [1, 1.0, 1.0, 64, 128, 200, 5, 1, 0.4, 100, 60, 600, 1, 1],
])


# ----------------------------------------------------------------------
# 8. __tables__.xlsx 登记新表
# ----------------------------------------------------------------------

print('[8/8] __tables__.xlsx（登记新表）')
p = os.path.join(DATAS, '__tables__.xlsx')
wb = openpyxl.load_workbook(p)
ws = wb['Sheet1']
existing = set()
for r in range(4, ws.max_row + 1):
    v = ws.cell(r, 2).value
    if v:
        existing.add(v)

new_rows = [
    ('子弹', 'TBBulletData', 'BulletData', 'True', 'BulletData.xlsx'),
    ('棋盘布局（替代原 ASCII 地图文件）', 'TBLevelMap', 'LevelMap', 'True', 'LevelMap.xlsx'),
    ('全局参数', 'TBGlobal', 'Global', 'True', 'Global.xlsx'),
]
r = ws.max_row + 1
for row in new_rows:
    if row[1] in existing:
        print('  已存在，跳过：' + row[1])
        continue
    for i, v in enumerate(row):
        ws.cell(r, i + 1).value = v
    print('  + %s (%s)' % (row[1], row[4]))
    r += 1
wb.save(p)

print()
print('全部完成。')

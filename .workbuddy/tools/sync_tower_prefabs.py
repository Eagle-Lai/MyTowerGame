# -*- coding: utf-8 -*-
"""
同步 Tower 目录：把 Power / Retard 的**分等级 prefab**（Tower_Power0/1/2、
Tower_Retard0/1/2）接入配置表与资源表。

背景
----
美术/策划新交付了带等级差异的塔模型：
  Assets/Prefabs/Tower/Power/Tower_Power0|1|2.prefab
  Assets/Prefabs/Tower/Retard/Tower_Retard0|1|2.prefab
每个 prefab 用的是对应等级的贴图（tower_base[_lv2|_lv3] 等）。

但配置表 tbtowerinfo.json 里 4~9 行仍指向旧的单文件
  Tower_Power / Tower_Retard
且 ResTable 里根本没有登记这两个逻辑名 —— 塔建不出来。

本脚本做的事（幂等，可重复执行）
--------------------------------
1. 改 Assets/ConfigJson/tbtowerinfo.json：
     id 4/5/6  ResName: Tower_Power  -> Tower_Power0/1/2
     id 7/8/9  ResName: Tower_Retard -> Tower_Retard0/1/2
2. 改 Luban 源表 Luban/Config/Datas/TowerInfo.xlsx 的同名列（按列名定位，不写死列号）
3. 输出 ResTable 需要新增的登记行（供人工/后续脚本核对）

用法
----
  python sync_tower_prefabs.py            # 预览（dry-run）
  python sync_tower_prefabs.py --apply    # 真正写入（先自动备份）
"""

import argparse
import os
import shutil
import sys
import json

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
JSON_PATH = os.path.join(ROOT, "Assets", "ConfigJson", "tbtowerinfo.json")
XLSX_PATH = os.path.join(ROOT, "Luban", "Config", "Datas", "TowerInfo.xlsx")

# id -> 新 ResName
RESNAME_PATCH = {
    4: "Tower_Power0",
    5: "Tower_Power1",
    6: "Tower_Power2",
    7: "Tower_Retard0",
    8: "Tower_Retard1",
    9: "Tower_Retard2",
}


def patch_json(apply):
    with open(JSON_PATH, "r", encoding="utf-8-sig") as f:
        rows = json.load(f)

    changed = []
    for r in rows:
        new = RESNAME_PATCH.get(r.get("id"))
        if new and r.get("resName") != new:
            changed.append((r["id"], r["name"], r["resName"], new))
            if apply:
                r["resName"] = new

    if not changed:
        print("[JSON] 无需修改（已是目标状态）")
        return

    for cid, name, old, new in changed:
        print("[JSON] id=%-2d %-12s %s -> %s" % (cid, name, old, new))

    if apply:
        shutil.copy2(JSON_PATH, JSON_PATH + ".bak")
        with open(JSON_PATH, "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2)
        print("[JSON] 已写入（备份 %s.bak）" % os.path.basename(JSON_PATH))
    else:
        print("[JSON] 预览模式，未写入（加 --apply 生效）")


def patch_xlsx(apply):
    try:
        import openpyxl
    except ImportError:
        print("[XLSX] 跳过：未安装 openpyxl")
        return

    if not os.path.exists(XLSX_PATH):
        print("[XLSX] 跳过：找不到 %s" % XLSX_PATH)
        return

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    # ---- 定位表头行：找到同时含 id 与 resName 的行 ----
    header_row = None
    col_id = col_res = None
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                key = v.strip().lstrip("#*&").strip()
                if key.lower() == "id":
                    header_row, col_id = r, c
                if key.lower() == "resname":
                    col_res = c
        if header_row and col_res:
            break

    if not (header_row and col_id and col_res):
        print("[XLSX] 找不到 id / resName 表头，跳过（请手工核对表结构）")
        return

    print("[XLSX] 表头行=%d, id 列=%d, resName 列=%d" % (header_row, col_id, col_res))

    changed = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=col_id).value
        if raw is None:
            continue
        try:
            rid = int(raw)
        except (TypeError, ValueError):
            continue
        new = RESNAME_PATCH.get(rid)
        if not new:
            continue
        cur = ws.cell(row=r, column=col_res).value
        if cur != new:
            changed.append((rid, r, cur, new))
            if apply:
                ws.cell(row=r, column=col_res).value = new

    if not changed:
        print("[XLSX] 无需修改（已是目标状态）")
        return

    for rid, r, old, new in changed:
        print("[XLSX] 行%d id=%-2d %s -> %s" % (r, rid, old, new))

    if apply:
        backup = XLSX_PATH + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(XLSX_PATH, backup)
        wb.save(XLSX_PATH)
        print("[XLSX] 已写入（备份 %s.bak）" % os.path.basename(XLSX_PATH))
    else:
        print("[XLSX] 预览模式，未写入（加 --apply 生效）")


def print_restable_hint():
    print()
    print("=" * 72)
    print("ResTable.cs 需要新增的登记行（供核对）")
    print("=" * 72)
    for rid in sorted(RESNAME_PATCH):
        name = RESNAME_PATCH[rid]
        bundle = "TowerPower" if "Power" in name else "TowerRetard"
        print('    { "%s", new ResAddress(ResBundle.%s, "%s", '
              '"Assets/Prefabs/Tower/%s/%s.prefab") },'
              % (name, bundle, name, "Power" if "Power" in name else "Retard", name))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="真正写入（默认只预览）")
    args = ap.parse_args()

    print("目标工程：%s" % ROOT)
    print("-" * 72)
    patch_json(args.apply)
    print("-" * 72)
    patch_xlsx(args.apply)
    print_restable_hint()


if __name__ == "__main__":
    main()

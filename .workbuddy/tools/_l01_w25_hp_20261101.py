#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""一次性手工改表：L01 最后一波（W25 = roundId 405）血量定向调整 —— 2026-11-01 11:13 用户口径。

    空中单位（isFlying=1）×1.03；地面单位（isFlying=0）×0.95。

基线 = 改动前的现值（乘性改动，重复执行会复利，故带幂等护栏：EnemyList.desc 含 TAG 即拒绝重跑）。
只动 W25 引用的 6 个变体行（地面 5 + 空中 1），其他波次、其他关卡、其他字段一律不动。
同时把 EnemyList 该行的行首 `EHP=` 汇总与追加 tag 一起更新，避免备注与实际脱节。

用法：
  python _l01_w25_hp_20261101.py values | apply | verify
"""
import json
import re
import pathlib
import shutil
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
ROUND_ID = 405          # L01 W25
GROUND = 0.95
AIR = 1.03
TAG = 'w25hp2-20261101'
BAK_ED = L01 / 'EnemyData.xlsx.bak_20261101_w25hp'
BAK_EL = L01 / 'EnemyList.xlsx.bak_20261101_w25hp'


def heads(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def w25_ids():
    wb = openpyxl.load_workbook(L01 / 'EnemyList.xlsx', data_only=True)
    ws = wb['Sheet1']
    h = heads(ws)
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, h['id']).value == ROUND_ID:
            ids = [int(x) for x in str(ws.cell(r, h['EnemyIndexs']).value).split(',') if str(x).strip()]
            desc = str(ws.cell(r, h['desc']).value or '')
            row = r
            wb.close()
            return sorted(set(ids)), ids, row, desc
    wb.close()
    raise SystemExit('找不到 W25 行')


def load_enemy(path):
    tmp = ROOT / '.workbuddy/tools/__pycache__/bakcheck'
    tmp.mkdir(parents=True, exist_ok=True)
    src = path
    if path.suffix != '.xlsx':
        src = tmp / 'EnemyData_bak.xlsx'
        shutil.copy2(path, src)
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[wb.sheetnames[0]]
    h = heads(ws)
    out = {}
    for r in range(4, ws.max_row + 1):
        eid = ws.cell(r, h['id']).value
        if eid is None:
            continue
        out[int(eid)] = dict(row=r, name=ws.cell(r, h['name']).value,
                            hp=float(ws.cell(r, h['hp']).value),
                            armor=float(ws.cell(r, h['armor']).value or 0),
                            fly=int(ws.cell(r, h['isFlying']).value or 0))
    wb.close()
    return out


def main():
    mode = sys.argv[1] if len(sys.argv) > 1 else 'values'
    ids, seq, el_row, desc = w25_ids()
    cur = load_enemy(L01 / 'EnemyData.xlsx')
    base = load_enemy(BAK_ED) if BAK_ED.exists() else cur

    plan = {}
    for i in ids:
        e = base[i]
        f = AIR if e['fly'] else GROUND
        plan[i] = (e['hp'], round(e['hp'] * f, 6), f)

    def ehp(d, seq_):
        return sum(d[i]['hp'] / (1 - d[i]['armor']) for i in seq_)

    old_ehp = ehp(base, seq)
    new_map = {i: dict(base[i], hp=plan[i][1]) for i in ids}
    merged = {**base, **new_map}
    new_ehp = ehp(merged, seq)

    if mode in ('values', 'verify'):
        print('id | 怪 | 飞行 | 旧 hp | 新 hp | 倍率')
        for i in ids:
            e = base[i]
            tag = '空中' if e['fly'] else '地面'
            print(f"{i} | {e['name']:<9} | {tag} | {e['hp']:.6f} | {plan[i][1]:.6f} | x{plan[i][2]}")
        print(f'EHP: {old_ehp:,.6f} -> {new_ehp:,.6f} ({100 * (new_ehp / old_ehp - 1):+.3f}%)')
        if mode == 'verify':
            bad = [i for i in ids if abs(cur[i]['hp'] - plan[i][1]) > 1e-6]
            if bad:
                print('❌ 校验失败：', bad)
                return 1
            print(f'✅ 校验通过：W25 共 {len(ids)} 行（空中 1 / 地面 {len(ids) - 1}）均符合预期')
        return 0

    assert mode == 'apply', f'未知模式 {mode}'
    if TAG in desc:
        print(f'⏭ 该波已带标记 {TAG}，拒绝重复叠加（乘性改动不可重跑）')
        return 0

    for src, dst in ((L01 / 'EnemyData.xlsx', BAK_ED), (L01 / 'EnemyList.xlsx', BAK_EL)):
        if not dst.exists():
            shutil.copy2(src, dst)
            print(f'备份 -> {dst.name}')

    wb = openpyxl.load_workbook(L01 / 'EnemyData.xlsx')
    ws = wb[wb.sheetnames[0]]
    h = heads(ws)
    for i in ids:
        ws.cell(cur[i]['row'], h['hp']).value = plan[i][1]
    wb.save(L01 / 'EnemyData.xlsx')

    # 备注：更新行首 EHP 汇总 + 追加 tag
    wb = openpyxl.load_workbook(L01 / 'EnemyList.xlsx')
    ws = wb['Sheet1']
    h = heads(ws)
    parts = desc.split(' | ')
    seg_new = re.sub(r'EHP=[\d.]+', f'EHP={new_ehp:.6f}', parts[0], count=1)
    rest = ' | '.join(parts[1:])
    new_desc = seg_new + (' | ' + rest if rest else '') + f" | {TAG}: 地面 x{GROUND}, 空中 x{AIR} (EHP={new_ehp:.6f})"
    ws.cell(el_row, h['desc']).value = new_desc
    wb.save(L01 / 'EnemyList.xlsx')

    print('已写入：')
    for i in ids:
        e = base[i]
        print(f"  {i} {e['name']:<9} {'空中' if e['fly'] else '地面'}  {e['hp']:.6f} -> {plan[i][1]:.6f}")
    print(f'EHP {old_ehp:,.6f} -> {new_ehp:,.6f} ({100 * (new_ehp / old_ehp - 1):+.3f}%)')
    print(f'备注已追加 {TAG}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

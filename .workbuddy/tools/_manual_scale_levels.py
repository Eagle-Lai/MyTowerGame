#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""一次性手工填表：把 L02~L30 按「第一关的数值关系」重算后写进各关 xlsx。

第一关（已人工调过）是基准：
  · 地面怪 hp = 生成值 ×1.5，飞行怪 hp = 生成值 ×2.0
  · 空军 W1~W4 不出，W5~W25 数量 ×2
  · 开局 100 金、12 命、25 波、每 1 点 EHP 换约 1/50 金币

铺到 L02~L30 的规则：
  k(L) = 1 + 0.2×(L-1)                 # 与开局金币 100+20(L-1) 同步
  目标 EHP(L)  = EHP(L1)  × k(L)
  目标 收入(L) = 收入(L1) × k(L)
  hp：先按角色 ×1.5 / ×2.0，再整关等比缩放到目标 EHP
  reward：按剩余额度整体缩放，保证「EHP/金币」= 第一关的 50.1
  波次奖励：× k(L)
  L01 一个字节都不动。

只改 xlsx；写完必须手工跑 Luban 导出 + 重封 runtimeContentHash。
"""
from __future__ import annotations

import argparse
import collections
import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
LEVELS = ROOT / 'Luban/Config/Datas/Levels'

K_SLOPE = 0.20
GROUND_HP = 1.5
AIR_HP = 2.0
AIR_FIRST_WAVE = 5
AIR_FACTOR = 2


def load(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    return wb, ws, heads


def read_level(L):
    wb_e, ws_e, h_e = load(LEVELS / f'L{L:02d}' / 'EnemyData.xlsx')
    enemies = {}
    for r in range(4, ws_e.max_row + 1):
        eid = ws_e.cell(r, h_e['id']).value
        if eid is None:
            continue
        desc = str(ws_e.cell(r, h_e['desc']).value or '')
        cid = re.search(r'canonical=(\d+)', desc)
        role = re.search(r'role=(\w+)', desc)
        enemies[int(eid)] = {
            'row': r, 'hp': float(ws_e.cell(r, h_e['hp']).value),
            'armor': float(ws_e.cell(r, h_e['armor']).value or 0),
            'reward': int(ws_e.cell(r, h_e['reward']).value or 0),
            'flying': int(ws_e.cell(r, h_e['isFlying']).value or 0),
            'canonical': int(cid.group(1)) if cid else 0,
            'role': role.group(1) if role else 'normal',
        }
    wb_l, ws_l, h_l = load(LEVELS / f'L{L:02d}' / 'EnemyList.xlsx')
    waves = {}
    for r in range(4, ws_l.max_row + 1):
        lid = ws_l.cell(r, h_l['id']).value
        if lid is None:
            continue
        desc = str(ws_l.cell(r, h_l['desc']).value or '')
        wave = int(re.match(r'L\d+ W(\d+)/', desc).group(1))
        ids = [int(x) for x in str(ws_l.cell(r, h_l['EnemyIndexs']).value).split(',') if str(x).strip()]
        waves[wave] = {'row': r, 'ids': ids}
    wb_r, ws_r, h_r = load(LEVELS / f'L{L:02d}' / 'RoundData.xlsx')
    rounds = {}
    for r in range(4, ws_r.max_row + 1):
        rid = ws_r.cell(r, h_r['id']).value
        if rid is None:
            continue
        rounds[rid] = {'row': r, 'rewardGold': int(ws_r.cell(r, h_r['rewardGold']).value or 0)}
    return ((wb_e, ws_e, h_e, enemies), (wb_l, ws_l, h_l, waves), (wb_r, ws_r, h_r, rounds))


def ehp_of(ids, enemies):
    return sum(enemies[i]['hp'] / (1 - enemies[i]['armor']) for i in ids)


def apply_air_rule(ids, enemies, wave):
    air = [i for i in ids if enemies[i]['flying']]
    ground = [i for i in ids if not enemies[i]['flying']]
    if not ground:
        return list(ids)
    if wave < AIR_FIRST_WAVE:
        top = collections.Counter(ground).most_common(1)[0][0]
        return ground + [top] * len(air)
    need = len(air) * AIR_FACTOR - len(air)
    if need <= 0:
        return list(ids)
    out = list(ids)
    pool = air or ground
    order = sorted((idx for idx, i in enumerate(out) if not enemies[i]['flying']),
                   key=lambda idx: enemies[out[idx]]['hp'])
    for k, idx in enumerate(order[:need]):
        out[idx] = pool[k % len(pool)]
    return out


def build_desc(L, wave, ids, enemies):
    grouped = collections.Counter((enemies[i]['role'], enemies[i]['canonical']) for i in ids)
    parts = ','.join(f'{role}:{cid}x{n}' for (role, cid), n in sorted(grouped.items()))
    return f'L{L} W{wave}/25; N={len(ids)}; EHP={ehp_of(ids, enemies):.6f}; {parts} | manual: L1-scaled'


def main(argv=None):
    parser = argparse.ArgumentParser(description='按第一关的数值关系重算 L02~L30')
    parser.add_argument('--slope', type=float, default=K_SLOPE)
    args = parser.parse_args(argv)
    k = lambda L: 1 + args.slope * (L - 1)

    (_, _, _, e1), (_, _, _, w1), (_, _, _, r1) = read_level(1)
    base_ids = [i for wave in sorted(w1) for i in w1[wave]['ids']]
    base_ehp = ehp_of(base_ids, e1)
    base_gold = sum(e1[i]['reward'] for i in base_ids) + sum(x['rewardGold'] for x in r1.values())
    print(f'基准 L1：n={len(base_ids)} air={sum(1 for i in base_ids if e1[i]["flying"])} '
          f'EHP={base_ehp:.0f} 收入={base_gold} EHP/金={base_ehp/base_gold:.1f}')

    rows = []
    for L in range(2, 31):
        (wb_e, ws_e, h_e, enemies), (wb_l, ws_l, h_l, waves), (wb_r, ws_r, h_r, rounds) = read_level(L)
        new_ids = {w: apply_air_rule(waves[w]['ids'], enemies, w) for w in waves}

        # 1) 血量：角色倍率 → 整关等比缩放到目标 EHP
        for e in enemies.values():
            e['hp'] *= AIR_HP if e['flying'] else GROUND_HP
        cur_ehp = sum(ehp_of(new_ids[w], enemies) for w in new_ids)
        s_hp = base_ehp * k(L) / cur_ehp

        # 2) 金币：波次奖励 ×k；击杀金币按剩余额度整体缩放，使 EHP/金 与 L1 一致
        for row in rounds.values():
            row['rewardGold'] = max(1, round(row['rewardGold'] * k(L)))
        wave_gold = sum(x['rewardGold'] for x in rounds.values())
        kill_now = sum(enemies[i]['reward'] for w in new_ids for i in new_ids[w])
        target_gold = base_gold * k(L)
        s_gold = max(1e-6, (target_gold - wave_gold) / kill_now)

        for e in enemies.values():
            ws_e.cell(e['row'], h_e['hp']).value = round(e['hp'] * s_hp, 6)
            e['hp'] *= s_hp
            e['reward'] = max(1, round(e['reward'] * s_gold))
            ws_e.cell(e['row'], h_e['reward']).value = e['reward']
        wb_e.save(LEVELS / f'L{L:02d}' / 'EnemyData.xlsx')

        for w in sorted(new_ids):
            ids = new_ids[w]
            ws_l.cell(waves[w]['row'], h_l['EnemyIndexs']).value = ','.join(str(x) for x in ids)
            ws_l.cell(waves[w]['row'], h_l['desc']).value = build_desc(L, w, ids, enemies)
        wb_l.save(LEVELS / f'L{L:02d}' / 'EnemyList.xlsx')

        for row in rounds.values():
            ws_r.cell(row['row'], h_r['rewardGold']).value = row['rewardGold']
        wb_r.save(LEVELS / f'L{L:02d}' / 'RoundData.xlsx')

        total_ehp = sum(ehp_of(new_ids[w], enemies) for w in new_ids)
        total_gold = sum(enemies[i]['reward'] for w in new_ids for i in new_ids[w]) + \
            sum(x['rewardGold'] for x in rounds.values())
        rows.append({'L': L, 'k': k(L), 'n': sum(len(v) for v in new_ids.values()),
                     'air': sum(1 for w in new_ids for i in new_ids[w] if enemies[i]['flying']),
                     'ehp': total_ehp, 'gold': total_gold, 'ratio': total_ehp / total_gold})

    print('\n L    k      数量   空军   EHP         收入    EHP/金')
    for r in rows:
        print(f"{r['L']:>2}  {r['k']:>5.2f}  {r['n']:>5}  {r['air']:>4}  {r['ehp']:>11.0f}  {r['gold']:>6}  {r['ratio']:>6.1f}")


if __name__ == '__main__':
    main()

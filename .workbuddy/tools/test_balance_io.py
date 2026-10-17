"""Real workbook staging and CLI process checks; formal project data stays read-only."""
import contextlib
import copy
import io
import json
import os
import pathlib
import subprocess
import sys
import tempfile
import time
import unittest
from unittest import mock

import openpyxl
import balance_pipeline as core


class BalanceIoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.formal_hashes={str(p):core.file_hash(p) for parent in [core.DATA,core.ROOT/'Assets/Gen',core.ROOT/'Assets/ConfigJson']
                           for p in parent.glob('*') if p.is_file()}
        cls.evidence={'tests':[]}

    @classmethod
    def tearDownClass(cls):
        after={p:core.file_hash(p) for p in cls.formal_hashes}
        assert after==cls.formal_hashes,'Formal project data changed during read-only staging tests'
        cls.evidence['formalProjectHashesUnchanged']=True
        core.write_json(core.ROOT/'Docs/BalanceEvidence/pipeline-io-tests.json',cls.evidence)

    def cli(self,command,out,editor_log=False,extra=()):
        env=os.environ.copy();env.pop('FT_BALANCE_EDITOR_LOG',None)
        if editor_log:env['FT_BALANCE_EDITOR_LOG']=str(out/'pipeline.log')
        started=time.perf_counter()
        result=subprocess.run([sys.executable,str(pathlib.Path(core.__file__)),command,'--out',str(out),*extra],
                              cwd=core.ROOT,env=env,capture_output=True,text=True,encoding='utf-8',timeout=30)
        receipt=core.read_json(out/'cli-exit.json')
        self.assertEqual(result.returncode,receipt['exitCode']);self.assertEqual(receipt['jobId'],out.name)
        self.assertEqual(receipt['command'],command);self.assertTrue(receipt['finishedUtc'].endswith('Z'))
        self.evidence['tests'].append({'name':self._testMethodName,'command':command,'exitCode':result.returncode,
                                      'elapsedSec':time.perf_counter()-started,'editorFileLog':editor_log})
        return result,receipt

    def test_real_cli_success_keeps_normal_stdout(self):
        with tempfile.TemporaryDirectory() as td:
            result,_=self.cli('inspect',pathlib.Path(td)/'success')
            self.assertEqual(result.returncode,0);self.assertIn('"stage": "complete"',result.stdout)

    def test_real_cli_editor_file_log_redirects_all_output_and_traceback(self):
        with tempfile.TemporaryDirectory() as td:
            out=pathlib.Path(td)/'editor-success';result,_=self.cli('inspect',out,True)
            self.assertEqual(result.returncode,0);self.assertEqual(result.stdout+result.stderr,'')
            self.assertIn('"stage": "complete"',(out/'pipeline.log').read_text(encoding='utf-8'))
            out=pathlib.Path(td)/'editor-failure'
            result,receipt=self.cli('inspect',out,True,['--recipe',str(pathlib.Path(td)/'missing.json')])
            self.assertEqual(result.returncode,1);self.assertEqual(result.stdout+result.stderr,'')
            self.assertIn('Traceback',(out/'pipeline.log').read_text(encoding='utf-8'));self.assertIn('error',receipt)

    def test_real_cli_argparse_failure_replaces_stale_success_receipt(self):
        with tempfile.TemporaryDirectory() as td:
            out=pathlib.Path(td)/'stale';core.write_json(out/'cli-exit.json',{'exitCode':0,'stale':True})
            result,receipt=self.cli('unsupported-command',out,True)
            self.assertEqual(result.returncode,2);self.assertNotIn('stale',receipt)
            self.assertIn('invalid choice',(out/'pipeline.log').read_text(encoding='utf-8'))

    def test_real_workbooks_reuse_rows_preserve_prototypes_and_export(self):
        recipe=core.read_json(core.ROOT/'Library/BalanceWorkbench/gradient-calibration/recipe.json')
        plans=[core.generate(recipe,p) for p in ['D08','D08','D04']]
        plans[1]=copy.deepcopy(plans[1]);plans[1]['tables']['tbenemydata']=plans[1]['tables']['tbenemydata'][:-64]
        original_data=core.DATA
        with tempfile.TemporaryDirectory(prefix='FT-balance-io-') as td:
            source=pathlib.Path(td)/'source';source.mkdir()
            import shutil
            for p in original_data.glob('*.xlsx'):shutil.copy2(p,source/p.name)
            book=openpyxl.load_workbook(source/'EnemyData.xlsx');sheet=book.active
            unknown_column=sheet.max_column+1;sheet.cell(1,unknown_column,'manualAnnotation')
            sheet.cell(2,unknown_column,'string');sheet.cell(3,unknown_column,'manual annotation')
            prototype_rows=[r for r in range(4,sheet.max_row+1) if isinstance(sheet.cell(r,2).value,int) and sheet.cell(r,2).value<100000]
            self.assertEqual(len(prototype_rows),132)
            for r in prototype_rows:sheet.cell(r,unknown_column,'preserve-'+str(sheet.cell(r,2).value))
            book.save(source/'EnemyData.xlsx');book.close()
            book=openpyxl.load_workbook(source/'EnemyData.xlsx');sheet=book.active
            prototypes={sheet.cell(r,2).value:tuple((sheet.cell(r,c).value,copy.copy(sheet.cell(r,c)._style))
                        for c in range(1,unknown_column+1)) for r in prototype_rows}
            book.close()
            row_counts=[];timings=[];events=[]
            try:
                for i,plan in enumerate(plans):
                    core.DATA=source;stage=pathlib.Path(td)/('stage-'+str(i));capture=io.StringIO();started=time.perf_counter()
                    with contextlib.redirect_stdout(capture):core.excel_write(stage,plan)
                    timings.append(time.perf_counter()-started)
                    events.extend(json.loads(line) for line in capture.getvalue().splitlines())
                    book=openpyxl.load_workbook(stage/'EnemyData.xlsx');sheet=book.active;row_counts.append(sheet.max_row)
                    ids={sheet.cell(r,2).value:r for r in range(4,sheet.max_row+1) if sheet.cell(r,2).value is not None}
                    self.assertEqual(len(ids),len(plan['tables']['tbenemydata']))
                    for rid,expected in prototypes.items():
                        r=ids[rid]
                        self.assertEqual(tuple((sheet.cell(r,c).value,sheet.cell(r,c)._style)
                                              for c in range(1,unknown_column+1)),expected)
                    book.close();source=stage
                self.assertEqual(len(set(row_counts)),1,'Switching/re-expanding profiles must reuse obsolete and blank variant rows')
                # Extra manual annotation remains a harmless unowned workbook column.
                code=pathlib.Path(td)/'Gen';data=pathlib.Path(td)/'ConfigJson';started=time.perf_counter()
                core.luban_export(source,code,data,pathlib.Path(td)/'luban-cache')
                content_hash=core.runtime_content_hash(plans[-1],data)
                for row in plans[-1]['tables']['tblevelrating']:row['runtimeContentHash']=content_hash
                core.write_rating_excel(source/'LevelRating.xlsx',plans[-1]['tables']['tblevelrating'])
                core.luban_export(source,code,data,pathlib.Path(td)/'luban-cache')
                export_sec=time.perf_counter()-started
                validation=core.validate(plans[-1],data)
                self.assertTrue(validation['passed'],validation['errors'])
                for p in code.glob('*.cs'):
                    raw=p.read_bytes();self.assertNotIn(b'\r\n',raw)
                    self.assertTrue(all(line==line.rstrip() for line in raw.splitlines()))
                self.evidence['tests'].append({'name':self._testMethodName,'stagingSeconds':timings,
                    'rowCounts':row_counts,'prototypeRowsPreserved':len(prototypes),'lubanExportSec':export_sec,'events':events})
            finally:core.DATA=original_data

    def test_excel_cancel_during_variant_rows_preserves_formal_data(self):
        plan=core.generate(core.read_json(core.DEFAULT_RECIPE),'D08');calls=0
        def cancel(stage):
            nonlocal calls
            calls+=1
            if calls==3:raise RuntimeError('injected cancellation during variant row staging')
        with tempfile.TemporaryDirectory() as td,mock.patch.object(core,'check_cancel',side_effect=cancel):
            with self.assertRaisesRegex(RuntimeError,'injected cancellation'):
                with contextlib.redirect_stdout(io.StringIO()):core.excel_write(pathlib.Path(td)/'stage',plan)
        self.evidence['tests'].append({'name':self._testMethodName,'safeBoundaryCalls':calls,'cancelled':True})


if __name__=='__main__':unittest.main(verbosity=2)

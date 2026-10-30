# -*- coding: utf-8 -*-
"""P3：把 Global 里的 14 个"玩法常量"下沉成每关一份 LevelRules（2026-10-30）。

拆分口径（用户已确认：引擎常量留全局）：
  留 Global.xlsx （7 个，引擎/性能）：
    enemyPoolSize, bulletPoolSize, maxEnemyAlive, targetFrameRate,
    deathRecycleDelayMs, checkOfflinePerFrame, showDebugLog
  下沉 Levels/Lxx/LevelRules.xlsx（14 个，玩法/表现）：
    cellSize, boardOriginPadding, defaultReward, defaultDamageToPlayer,
    defaultBodyRadius, defaultSearchIntervalMs, autoNextRoundDelayMs,
    autoStartRoundDelayMs, sellRefundRate, artFaceLeft, showDamageText,
    damageTextThrottleMs, shakeAmplitude, shakeDurationMs

用法：
  python gen_level_rules.py [--write]
"""
import pathlib
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
LEVELS = DATAS / 'Levels'

ENGINE_KEYS = ['enemyPoolSize', 'bulletPoolSize', 'maxEnemyAlive', 'targetFrameRate',
               'deathRecycleDelayMs', 'checkOfflinePerFrame', 'showDebugLog']
RULE_KEYS = ['cellSize', 'boardOriginPadding', 'defaultReward', 'defaultDamageToPlayer',
             'defaultBodyRadius', 'defaultSearchIntervalMs', 'autoNextRoundDelayMs',
             'autoStartRoundDelayMs', 'sellRefundRate', 'artFaceLeft', 'showDamageText',
             'damageTextThrottleMs', 'shakeAmplitude', 'shakeDurationMs']


def main(write):
    p = DATAS / 'Global.xlsx'
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    heads = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v:
            heads[v] = c
    missing = [k for k in RULE_KEYS + ENGINE_KEYS if k not in heads]
    if missing:
        raise RuntimeError('Global.xlsx 缺列：%s' % missing)

    header = {k: [ws.cell(r, heads[k]).value for r in (1, 2, 3)] for k in RULE_KEYS}
    values = {k: ws.cell(4, heads[k]).value for k in RULE_KEYS}
    print('从 Global 读到 14 个玩法常量：cellSize=%s sellRefundRate=%s autoNextRoundDelayMs=%s'
          % (values['cellSize'], values['sellRefundRate'], values['autoNextRoundDelayMs']))

    # 1) 生成 30 份 LevelRules.xlsx
    for lv in range(1, 31):
        folder = LEVELS / ('L%02d' % lv)
        if not folder.exists():
            raise RuntimeError('缺目录 %s' % folder)
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = 'Sheet1'
        ws2.cell(1, 1, '##')
        ws2.cell(2, 1, '##type')
        ws2.cell(3, 1, '##')
        ws2.cell(1, 2, 'id')
        ws2.cell(2, 2, 'int')
        ws2.cell(3, 2, '关卡 id（本表每关一份，2026-10-30 从 Global 拆分）')
        for i, k in enumerate(RULE_KEYS):
            c = 3 + i
            ws2.cell(1, c, k)
            ws2.cell(2, c, header[k][1])
            ws2.cell(3, c, header[k][2])
            ws2.cell(4, c, values[k])
            ws2.column_dimensions[get_column_letter(c)].width = 16
        ws2.cell(4, 2, lv)
        ws2.column_dimensions['A'].width = 6
        ws2.column_dimensions['B'].width = 8
        if write:
            wb2.save(folder / 'LevelRules.xlsx')
        wb2.close()
    print('%s：生成 30 份 Levels/Lxx/LevelRules.xlsx' % ('已写入' if write else 'dry-run'))

    # 2) 裁剪 Global.xlsx：只留 id + 7 个引擎常量
    if write:
        for k in sorted(RULE_KEYS, key=lambda x: -heads[x]):
            ws.delete_cols(heads[k])
        wb.save(p)
    wb.close()
    print('%s：Global.xlsx 删除 14 列玩法常量，只保留 id + 7 个引擎常量' % ('已写入' if write else 'dry-run'))

    # 3) 登记 TBLevelRules
    tp = DATAS / '__tables__.xlsx'
    wb3 = openpyxl.load_workbook(tp)
    ws3 = wb3.worksheets[0]
    h3 = {ws3.cell(1, c).value: c for c in range(1, ws3.max_column + 1) if ws3.cell(1, c).value}
    exists = any(ws3.cell(r, h3['full_name']).value == 'TBLevelRules'
                 for r in range(4, ws3.max_row + 1))
    if not exists:
        r = ws3.max_row + 1
        ws3.cell(r, h3['##var'], '每关玩法常数（2026-10-30 从 Global 拆出）')
        ws3.cell(r, h3['full_name'], 'TBLevelRules')
        ws3.cell(r, h3['value_type'], 'LevelRules')
        ws3.cell(r, h3['define_from_file'], 'True')
        ws3.cell(r, h3['input'], ','.join('Levels/L%02d/LevelRules.xlsx' % lv for lv in range(1, 31)))
        print('%s：__tables__.xlsx 新增 TBLevelRules（30 个文件）' % ('已写入' if write else 'dry-run'))
    else:
        print('TBLevelRules 已登记，跳过')
    if write:
        wb3.save(tp)
    wb3.close()


if __name__ == '__main__':
    main('--write' in sys.argv)

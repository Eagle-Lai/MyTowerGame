# -*- coding: utf-8 -*-
"""P5：怪物表彻底按关归属（2026-10-30）。

背景：
  `TBEnemyData` 现在是 31 个输入文件 = 基表 `EnemyData.xlsx`（132 行 canonical，id 1..132）
  + 30 个分层变体文件。实测：**没有任何波次引用 id < 1000 的行**（最小被引用 id = 3140136），
  基表纯属历史残留，却让"怪物配置"看起来像有一张公共表。

做法：
  1. 给 30 个 `Levels/Lxx/EnemyData.xlsx` 末尾加 `levelId`（int），值 = 该关 id；
  2. 把基表 `EnemyData.xlsx` 从 `__tables__.xlsx` 的 input 里摘掉（**文件保留**，作历史参照，
     类似旧 TowerInfo/BulletData）；
  3. 导出后 TBEnemyData = 30 个文件、12,462 行，全部带 levelId。

用法：
  python split_enemy_tables.py [--write]
"""
import pathlib
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
LEVELS = DATAS / 'Levels'


def main(write):
    # 1) 给每关的 EnemyData 加 levelId 列
    for lv in range(1, 31):
        p = LEVELS / ('L%02d' % lv) / 'EnemyData.xlsx'
        wb = openpyxl.load_workbook(p)
        ws = wb.worksheets[0]
        heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
        if 'levelId' in heads:
            print('L%02d 已有 levelId，跳过' % lv)
            wb.close()
            continue
        col = ws.max_column + 1
        ws.cell(1, col, 'levelId')
        ws.cell(2, col, 'int')
        ws.cell(3, col, '所属关卡 id（本表每关独立一份，2026-10-30 拆分）')
        n = 0
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value is not None:
                ws.cell(r, col, lv)
                n += 1
        ws.column_dimensions[get_column_letter(col)].width = 10
        if write:
            wb.save(p)
        wb.close()
        if lv in (1, 30):
            print('L%02d 写入 levelId=%d（%d 行）' % (lv, lv, n))
    print('%s：30 份关卡怪物表加 levelId' % ('已写入' if write else 'dry-run'))

    # 2) 摘掉基表
    tp = DATAS / '__tables__.xlsx'
    wb = openpyxl.load_workbook(tp)
    ws = wb.worksheets[0]
    heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    changed = 0
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['full_name']).value == 'TBEnemyData':
            ws.cell(r, heads['input'], ','.join('Levels/L%02d/EnemyData.xlsx' % lv for lv in range(1, 31)))
            changed += 1
    if changed != 1:
        raise RuntimeError('__tables__ 里没找到 TBEnemyData（changed=%d）' % changed)
    if write:
        wb.save(tp)
    wb.close()
    print('%s：__tables__.xlsx 的 TBEnemyData input 改为 30 个关卡文件（基表不再参与导出）'
          % ('已写入' if write else 'dry-run'))


if __name__ == '__main__':
    main('--write' in sys.argv)

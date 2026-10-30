# -*- coding: utf-8 -*-
"""P0：把塔表/子弹表拆成"每关一套"（2026-10-30 用户需求）。

背景：
  当前 `TowerInfo.xlsx`（9 行）与 `BulletData.xlsx`（6 行）是全局表，导致"某关想改塔数值"
  必须改全局、影响 30 关。本脚本把这两张表按关复制成 30 份，并重编号让主键全局唯一。

ID 规则（写死，代码侧按同一规则解析）：
  塔：  id = levelId * 1000 + 原 id       （L01 → 1001..1009）
        upgradeTo / bulletId 同步重写成本关内 id，跨关不串
  子弹：id = levelId * 100  + 原 id       （L01 → 101,102,104..107）
  两张表都新增末列 `levelId`（int），显式标明归属，不靠 id 反推。

用法：
  python gen_per_level_tower_bullet.py [--write]
  不带 --write 只做 dry-run 打印。
"""
import pathlib
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
LEVELS = DATAS / 'Levels'
TOWER_ID_BASE = 1000
BULLET_ID_BASE = 100


def read_sheet(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    heads = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v and v != '##':
            heads[v] = c
    rows = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is None:
            continue
        rows.append({k: ws.cell(r, c).value for k, c in heads.items()})
    header = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in (1, 2, 3)]
    widths = {c: ws.column_dimensions[get_column_letter(c)].width for c in range(1, ws.max_column + 1)}
    wb.close()
    return header, rows, widths, ws.max_column


def write_sheet(path, header, rows, widths, extra_head, extra_type, extra_note):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    n = len(header[0]) + 1
    for r in range(3):
        for c in range(1, len(header[r]) + 1):
            ws.cell(r + 1, c, header[r][c - 1])
    ws.cell(1, n, extra_head)
    ws.cell(2, n, extra_type)
    ws.cell(3, n, extra_note)
    for i, row in enumerate(rows):
        r = 4 + i
        for c in range(1, len(header[0]) + 1):
            key = header[0][c - 1]
            if key and key != '##':
                ws.cell(r, c, row.get(key))
        ws.cell(r, n, row['levelId'])
    for c, w in widths.items():
        if w:
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.column_dimensions[get_column_letter(n)].width = 14
    wb.save(path)
    wb.close()


def main(write):
    tower_header, tower_rows, tower_w, _ = read_sheet(DATAS / 'TowerInfo.xlsx')
    bullet_header, bullet_rows, bullet_w, _ = read_sheet(DATAS / 'BulletData.xlsx')
    print('源表：塔 %d 行，子弹 %d 行' % (len(tower_rows), len(bullet_rows)))

    made = 0
    for lv in range(1, 31):
        folder = LEVELS / ('L%02d' % lv)
        if not folder.exists():
            raise RuntimeError('关卡目录不存在：%s' % folder)

        towers = []
        for r in tower_rows:
            t = dict(r)
            t['id'] = lv * TOWER_ID_BASE + int(r['id'])
            up = int(r['upgradeTo'] or 0)
            t['upgradeTo'] = lv * TOWER_ID_BASE + up if up > 0 else 0
            bi = int(r['bulletId'] or 0)
            t['bulletId'] = lv * BULLET_ID_BASE + bi if bi > 0 else 0
            t['levelId'] = lv
            towers.append(t)

        bullets = []
        for r in bullet_rows:
            b = dict(r)
            b['id'] = lv * BULLET_ID_BASE + int(r['id'])
            b['levelId'] = lv
            bullets.append(b)

        if write:
            write_sheet(folder / 'TowerInfo.xlsx', tower_header, towers, tower_w,
                        'levelId', 'int', '所属关卡 id（本表每关独立一份，2026-10-30 拆分）')
            write_sheet(folder / 'BulletData.xlsx', bullet_header, bullets, bullet_w,
                        'levelId', 'int', '所属关卡 id（本表每关独立一份，2026-10-30 拆分）')
        made += 1
        if lv in (1, 30):
            print('L%02d 塔 id: %s' % (lv, [t['id'] for t in towers]))
            print('L%02d 塔 upgradeTo: %s   bulletId: %s' % (lv, [t['upgradeTo'] for t in towers], [t['bulletId'] for t in towers]))
            print('L%02d 子弹 id: %s' % (lv, [b['id'] for b in bullets]))
    print('%s：生成 %d 个关卡目录的 TowerInfo/BulletData' % ('已写入' if write else 'dry-run', made))

    # 更新 __tables__.xlsx 的 input 列
    p = DATAS / '__tables__.xlsx'
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    changed = 0
    for r in range(4, ws.max_row + 1):
        full = ws.cell(r, heads['full_name']).value
        if full == 'TBTowerInfo':
            ws.cell(r, heads['input'], ','.join('Levels/L%02d/TowerInfo.xlsx' % lv for lv in range(1, 31)))
            changed += 1
        elif full == 'TBBulletData':
            ws.cell(r, heads['input'], ','.join('Levels/L%02d/BulletData.xlsx' % lv for lv in range(1, 31)))
            changed += 1
    if changed != 2:
        raise RuntimeError('__tables__ 里没找到两张目标表（changed=%d）' % changed)
    if write:
        wb.save(p)
    wb.close()
    print('%s：__tables__.xlsx 已把 TBTowerInfo / TBBulletData 的 input 改为 30 个关卡文件'
          % ('已写入' if write else 'dry-run'))


if __name__ == '__main__':
    main('--write' in sys.argv)

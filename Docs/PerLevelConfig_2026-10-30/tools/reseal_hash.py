# -*- coding: utf-8 -*-
"""统一重封 runtimeContentHash（2026-10-30 起为唯一权威实现）。

【为什么要有这个脚本】工程里散落着 5 份各自复制粘贴的 compute_hash（l01_flat_drop /
l01_w20_soften / l01_wavegold_air / l01_ground_plus5 / l01_ground_econ_neutral），
一旦哈希口径变化（例如新增 tblevelrules），漏改一份就会算出错误哈希、Play Mode 直接报
"content hash mismatch"。这里集中一份，其余脚本以后改为调用它。

口径必须与 `Configs.CalculateRuntimeContentHash` 逐字一致：
  1) 依次拼接 KEYS 里每张表的 JSON 文本（去 BOM、CRLF→LF），格式 `key\\n{json}\\n`；
  2) 再按 id 升序拼接每行评级字段（`|` 分隔，列表用 `,` 连接，不含 runtimeContentHash）；
  3) SHA256 → 小写 hex。

用法：
  python reseal_hash.py print     # 只打印当前算出的哈希
  python reseal_hash.py write     # 写进 30 关 LevelRating.xlsx
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
JSON_DIR = ROOT / 'Assets/ConfigJson'

# 必须与 Configs.CalculateRuntimeContentHash 的 keys 数组完全一致（顺序也要一致）
KEYS = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
        'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal', 'tblevelrules']


def compute_hash():
    text = []
    for k in KEYS:
        p = JSON_DIR / (k + '.json')
        s = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        text.append('%s\n%s\n' % (k, s))
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows.sort(key=lambda r: r['id'])
    for r in rows:
        def joined(v):
            return ','.join(str(x) for x in v) if isinstance(v, list) else ''
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  joined(r.get('bossVariantIds')), joined(r.get('criticalVariantIds')),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(x) for x in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()


def write_hash(h):
    for i in range(1, 31):
        p = DATAS / ('Levels/L%02d/LevelRating.xlsx' % i)
        wb = openpyxl.load_workbook(p)
        ws = wb.worksheets[0]
        heads = {c.value: c.column for c in ws[1] if c.value}
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value is not None:
                ws.cell(r, heads['runtimeContentHash'], h)
        wb.save(p)
        wb.close()


if __name__ == '__main__':
    mode = sys.argv[1] if len(sys.argv) > 1 else 'print'
    h = compute_hash()
    print(h)
    if mode == 'write':
        write_hash(h)
        print('已写入 30 关 LevelRating.xlsx')

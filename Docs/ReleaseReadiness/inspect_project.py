"""Read-only project inventory. Writes evidence only under this script's directory.

Run with a Python that has openpyxl installed: python inspect_project.py
No Unity generation, import, save migration or source mutation is performed.
"""
from __future__ import annotations
import collections
import datetime as dt
import hashlib
import json
import pathlib
import re
import subprocess
import zipfile

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = pathlib.Path(__file__).resolve().parent / "evidence"

def rel(p):
    return str(p.relative_to(ROOT)).replace("\\", "/")

def read(p):
    return p.read_text(encoding="utf-8-sig", errors="replace")

def git(*args):
    return subprocess.check_output(["git", "-c", "core.quotepath=false", *args], cwd=ROOT).decode("utf-8", "replace").strip()

def save(name, obj):
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    assets = list((ROOT / "Assets").rglob("*"))
    files = [p for p in assets if p.is_file()]
    inventory = {
        "timestamp": dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(),
        "head": git("rev-parse", "HEAD"),
        "status": git("status", "--short"),
        "unity": read(ROOT / "ProjectSettings/ProjectVersion.txt"),
        "counts_by_extension": dict(collections.Counter(p.suffix.lower() for p in files)),
        "source_hashes": {rel(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files if p.suffix in {".cs", ".unity", ".prefab", ".asmdef", ".json", ".xlsx"}},
        "area_counts": {},
    }
    for area in ["Assets/Scripts", "Assets/Editor", "Assets/AOTAssembly", "Assets/Prefabs/Enemy", "Assets/Prefabs/Tower", "Assets/Prefabs/UI", "Assets/Audio", "Assets/StreamingAssets", "Assets/_UIAssets/Monsters"]:
        ps = [p for p in (ROOT / area).rglob("*") if p.is_file() and p.suffix != ".meta"]
        inventory["area_counts"][area] = {"files": len(ps), "bytes": sum(p.stat().st_size for p in ps), "extensions": dict(collections.Counter(p.suffix for p in ps))}
    save("baseline.json", inventory)

    guid_paths = collections.defaultdict(list)
    orphan_meta = []
    for base in [ROOT / "Assets", ROOT / "Packages", ROOT / "Library/PackageCache"]:
        for p in base.rglob("*.meta"):
            m = re.search(r"^guid:\s*([a-fA-F0-9]{32})", read(p), re.M)
            if m:
                guid_paths[m[1].lower()].append(p.with_suffix(""))
            if base.name == "Assets" and not p.with_suffix("").exists():
                orphan_meta.append(rel(p))
    unresolved = []
    missing_script = []
    dependencies = {}
    yaml_ext = {".prefab", ".unity", ".asset", ".controller", ".overrideController", ".anim", ".mat", ".playable", ".renderTexture", ".spriteatlas"}
    scanned = 0
    for p in files:
        if p.suffix not in yaml_ext:
            continue
        txt = read(p)
        if not txt.startswith("%YAML"):
            continue
        scanned += 1
        refs = []
        for line, text in enumerate(txt.splitlines(), 1):
            for m in re.finditer(r"guid:\s*([a-fA-F0-9]{32})", text):
                g = m[1].lower()
                if g.startswith("0000000000000000"):
                    continue
                paths = guid_paths.get(g, [])
                refs.extend(rel(t) for t in paths if t.is_relative_to(ROOT / "Assets"))
                if not paths:
                    item = {"path": rel(p), "line": line, "guid": g, "text": text.strip()}
                    unresolved.append(item)
                    if "m_Script:" in text:
                        missing_script.append(item)
        dependencies[rel(p)] = sorted(set(refs))
    roots = [rel(p) for p in files if "/Resources/" in rel(p)] + ["Assets/Scenes/Boot.unity", "Assets/Scenes/Main.unity"]
    roots += [rel(p) for p in files if rel(p).startswith("Assets/Prefabs/") and p.suffix == ".prefab"]
    visited = set()
    queue = list(roots)
    while queue:
        p = queue.pop()
        if p in visited:
            continue
        visited.add(p)
        queue.extend(dependencies.get(p, []))
    for item in unresolved:
        item["reachable_from_audit_roots"] = item["path"] in visited
    save("asset_scan.json", {"scope": "Assets YAML; GUID index includes Packages and local PackageCache; roots include all production Prefabs, both scenes, Resources (conservative reachability)", "scanned_yaml": scanned, "orphan_meta": orphan_meta, "duplicate_asset_guids": {g: [rel(p) for p in ps] for g, ps in guid_paths.items() if len([p for p in ps if p.is_relative_to(ROOT / "Assets")]) > 1}, "missing_script_guids": missing_script, "unresolved_guids": unresolved, "reachable_files": len(visited)})

    tables = {p.stem: json.loads(read(p)) for p in (ROOT / "Assets/ConfigJson").glob("*.json")}
    errors = []
    by_id = {}
    for name, rows in tables.items():
        by_id[name] = {str(r.get("id", r.get("Id"))): r for r in rows}
        ids = [str(r.get("id", r.get("Id"))) for r in rows]
        for key, n in collections.Counter(ids).items():
            if n > 1:
                errors.append({"table": name, "id": key, "error": "duplicate id"})
    def fk(table, row, field, target, optional=False):
        value = row.get(field)
        vals = value if isinstance(value, list) else [value]
        for v in vals:
            if optional and (v == 0 or v is None):
                continue
            if str(v) not in by_id[target]:
                errors.append({"table": table, "id": row.get("id"), "field": field, "value": v, "target": target})
    for r in tables["tbsceneinfo"]:
        fk("tbsceneinfo", r, "mapId", "tblevelmap")
        fk("tbsceneinfo", r, "RoundList", "tbrounddata")
    for r in tables["tbtowerinfo"]:
        fk("tbtowerinfo", r, "upgradeTo", "tbtowerinfo", True)
        fk("tbtowerinfo", r, "bulletId", "tbbulletdata", True)
    for r in tables["tbrounddata"]:
        fk("tbrounddata", r, "EnemyIndexs", "tbenemylist")
    for r in tables["tbenemylist"]:
        fk("tbenemylist", r, "EnemyIndexs", "tbenemydata")
    maps = []
    for r in tables["tblevelmap"]:
        cells = r["cells"]
        starts = [(rr, cc) for rr, line in enumerate(cells) for cc, ch in enumerate(line) if ch == "S"]
        ends = [(rr, cc) for rr, line in enumerate(cells) for cc, ch in enumerate(line) if ch == "E"]
        seen = set(starts)
        q = collections.deque(starts)
        while q:
            rr, cc = q.popleft()
            for nr, nc in [(rr-1,cc),(rr+1,cc),(rr,cc-1),(rr,cc+1)]:
                if 0 <= nr < len(cells) and 0 <= nc < len(cells[nr]) and cells[nr][nc] in ".SEP" and (nr,nc) not in seen:
                    seen.add((nr,nc)); q.append((nr,nc))
        maps.append({"id": r["id"], "rows": r["rows"], "cols": r["cols"], "dimension_valid": len(cells)==r["rows"] and all(len(s)==r["cols"] for s in cells), "starts": starts, "ends": ends, "all_exits_reachable": bool(ends) and all(p in seen for p in ends)})
    save("config_scan.json", {"counts": {k: len(v) for k,v in tables.items()}, "foreign_key_errors": errors, "maps": maps, "towers": tables["tbtowerinfo"], "bullets": tables["tbbulletdata"], "global": tables["tbglobal"], "waves_sample": tables["tbenemylist"][:2], "rounds_sample": tables["tbrounddata"][:2]})

    try:
        import openpyxl
        sheet_summary = []
        for p in (ROOT / "Luban/Config/Datas").glob("*.xlsx"):
            if p.name.startswith("~$"):
                continue
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            for sheet in wb:
                rows = list(sheet.values)
                sheet_summary.append({"file": rel(p), "sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "rows": rows})
            wb.close()
        save("source_tables.json", sheet_summary)
    except ImportError as exc:
        save("source_tables.json", {"unavailable": str(exc)})

    apks = []
    for p in (ROOT / "Builds").rglob("*.apk"):
        with zipfile.ZipFile(p) as z:
            items = z.infolist()
            groups = collections.Counter()
            for item in items:
                parts = item.filename.split("/")
                groups["/".join(parts[:3])] += item.file_size
            apks.append({"path": rel(p), "bytes": p.stat().st_size, "modified": dt.datetime.fromtimestamp(p.stat().st_mtime,dt.timezone(dt.timedelta(hours=8))).isoformat(), "sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "uncompressed_bytes": sum(i.file_size for i in items), "top_entries": [{"name":i.filename,"bytes":i.file_size,"compressed":i.compress_size} for i in sorted(items,key=lambda i:i.file_size,reverse=True)[:25]], "groups": groups.most_common(25), "abis": sorted(set(i.filename.split('/')[1] for i in items if i.filename.startswith('lib/'))), "dll_files": [i.filename for i in items if i.filename.endswith('.dll') or i.filename.endswith('/md5.txt')], "version_files": {i.filename: z.read(i).decode('utf-8','replace') for i in items if i.filename.endswith('.version')}})
    save("apk_inventory.json", apks)
    print(json.dumps({"counts": inventory["area_counts"], "tables": {k:len(v) for k,v in tables.items()}, "fk_errors": len(errors), "missing_script": len(missing_script), "unresolved_guids":len(unresolved), "orphan_meta":len(orphan_meta), "apks": [{k:a[k] for k in ['path','bytes','modified']} for a in apks]},ensure_ascii=False,indent=2))

if __name__ == "__main__":
    main()

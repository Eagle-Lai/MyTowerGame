# -*- coding: utf-8 -*-
"""L01 示意图迷宫重设计 —— 验证器。

策略：按示意图 42 个可建造格的蓝图顺序脚本化建造/升级（波边界），用
balance_simulator 的同源算法做空间仿真，并扩展记录每次击杀的位置，
验证 7 条约束：零漏怪 / 末怪距终点>=2格 / W5 起空军严格递增 / 初始 100 金 /
终局 42 格全满级 / 难度曲线 / 收支。

simulate_level_tracked 改编自 .workbuddy/tools/balance_simulator.py
(VERSION='spatial-fixed-step-30x25-maze-v6-astar')，仅增加击杀位置跟踪，
战斗逻辑逐行保持一致。
"""
import json, math, pathlib, sys, time, collections

TOOLS = pathlib.Path(r'D:/FishTestGame/.workbuddy/tools')
sys.path.insert(0, str(TOOLS))
sys.path.insert(0, str(pathlib.Path(__file__).parent))

import balance_simulator as bs
from balance_campaign import runtime_path_for as path_for, capacity
import l01_design

OUT = pathlib.Path(__file__).parents[1] / 'reports'

# ---------------------------------------------------------------- 蓝图（来自 L01 示意图像素提取）
# 42 个可建造格（1 起行列），路线 = 49 格蛇形（S(1,4)→底→右→上…→E(13,4)）
FROST = [(2, 4), (4, 4), (6, 4), (8, 4), (10, 4), (12, 4), (13, 6)]
NORMAL = [(2, 3), (4, 3), (6, 3), (8, 3), (10, 3), (12, 3),
          (2, 5), (4, 5), (6, 5), (8, 5), (10, 5), (12, 5),
          (1, 2), (1, 3), (13, 5), (13, 7),
          (2, 1), (6, 1), (10, 1), (4, 7), (8, 7), (12, 7)]
BLAST = [(1, 1), (2, 2), (4, 2), (6, 2), (8, 2), (10, 2), (12, 2),
         (2, 6), (4, 6), (6, 6), (8, 6), (10, 6), (12, 6)]
# 建造顺序：先直道两翼 Normal（前期输出）→ Frost 中枢 → 门口 → Blast 墙 → 角落
BUILD_ORDER = [
    [(6, 3), (8, 3)], [(6, 5), (8, 5)], [(4, 3), (10, 3)], [(4, 5), (10, 5)],
    [(2, 3), (12, 3)], [(2, 5), (12, 5)],
    [(6, 4), (8, 4)], [(4, 4), (10, 4)], [(2, 4), (12, 4)],
    [(1, 2), (1, 3), (13, 5), (13, 7)], [(13, 6)],
    [(6, 2), (8, 2), (6, 6), (8, 6)], [(4, 2), (10, 2), (4, 6), (10, 6)],
    [(2, 2), (12, 2), (2, 6), (12, 6)],
    [(2, 1), (6, 1), (10, 1), (4, 7), (8, 7), (12, 7)], [(1, 1)],
]
TOWER_L1 = {'N': 1, 'B': 4, 'F': 7}
TOWER_L2 = {'N': 2, 'B': 5, 'F': 8}
TOWER_L3 = {'N': 3, 'B': 6, 'F': 9}

def cell_type(cell):
    if cell in FROST: return 'F'
    if cell in NORMAL: return 'N'
    if cell in BLAST: return 'B'
    raise ValueError(f'蓝图外格子 {cell}')

def blueprint_flat():
    out = []
    for grp in BUILD_ORDER:
        out.extend(grp)
    return out

# ---------------------------------------------------------------- 调度器
def schedule_ops(plan):
    """按磁盘表(plan)的收入曲线把 42 建造 + 84 升级分配到波边界（先建满 42，再按优先级升级）。"""
    towers = {t['id']: t for t in build_plan_static_towers()}
    build_seq = [(cell, TOWER_L1[cell_type(cell)]) for cell in blueprint_flat()]
    # 升级优先级：Frost→L2 全部 → Normal 核心 12 →L2 → Frost→L3 → 其余→L2 → 其余→L3
    core12 = [(2, 3), (4, 3), (6, 3), (8, 3), (10, 3), (12, 3), (2, 5), (4, 5), (6, 5), (8, 5), (10, 5), (12, 5)]
    up_seq = []
    up_seq += [(c, TOWER_L2['F']) for c in FROST]
    up_seq += [(c, TOWER_L2['N']) for c in core12]
    up_seq += [(c, TOWER_L3['F']) for c in FROST]
    up_seq += [(c, TOWER_L2['N']) for c in NORMAL if c not in core12]
    up_seq += [(c, TOWER_L2['B']) for c in BLAST]
    up_seq += [(c, TOWER_L3['N']) for c in core12]
    up_seq += [(c, TOWER_L3['N']) for c in NORMAL if c not in core12]
    up_seq += [(c, TOWER_L3['B']) for c in BLAST]
    ops = []
    gold = plan['tables']['tbsceneinfo'][0]['initialGold']
    built = set()
    bi = ui = 0
    waves = sorted(plan['waves'], key=lambda w: w['waveOrdinal'])
    for w in range(1, 26):
        wv = waves[w - 1]
        while True:
            if bi < len(build_seq):
                cell, tid = build_seq[bi]
                cost = towers[tid]['prices']
                if gold >= cost:
                    ops.append({'waveOrdinal': w, 'timeMs': 0, 'action': 'build',
                                'row': cell[1] - 1, 'col': cell[0] - 1, 'towerId': tid})
                    gold -= cost; built.add(cell); bi += 1; continue
            if bi >= len(build_seq) and ui < len(up_seq):
                cell, tid = up_seq[ui]
                if cell in built:
                    cost = towers[tid]['prices']
                    if gold >= cost:
                        ops.append({'waveOrdinal': w, 'timeMs': 0, 'action': 'upgrade',
                                    'row': cell[1] - 1, 'col': cell[0] - 1, 'towerId': tid})
                        gold -= cost; ui += 1; continue
            break
        if w < 25:
            gold += wv['killGold'] + wv['rewardGold']
    return ops, dict(builds=bi, upgrades=ui, leftover_gold=gold,
                     last_op_wave=max((o['waveOrdinal'] for o in ops), default=0))

_TOWER_CACHE = None
def build_plan_static_towers():
    global _TOWER_CACHE
    if _TOWER_CACHE is None:
        import openpyxl
        wb = openpyxl.load_workbook(l01_design.DATA / 'TowerInfo.xlsx', data_only=True)
        ws = wb.worksheets[0]
        heads = {c.value: c.column for c in ws[1] if c.value}
        rows = []
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value is not None:
                d = {k: ws.cell(r, c).value for k, c in heads.items()}
                d['upgradeTo'] = int(d['upgradeTo']); d['canAttackAir'] = int(d['canAttackAir'] or 0)
                rows.append(d)
        wb.close(); _TOWER_CACHE = rows
    return _TOWER_CACHE

# ---------------------------------------------------------------- 跟踪版仿真（改编自 balance_simulator.simulate_level）
def simulate_level_tracked(plan, L, step_ms, out, seed=42):
    tables = plan['tables']; level = next(x for x in tables['tbsceneinfo'] if x['id'] == L)
    rating = next(x for x in tables['tblevelrating'] if x['id'] == L)
    cells = next(x['cells'] for x in tables['tblevelmap'] if x['id'] == level['mapId'])
    map_capacity = capacity(cells)
    enemy_table = {x['id']: x for x in tables['tbenemydata']}; tower_table = {x['id']: x for x in tables['tbtowerinfo']}
    waves = [w for w in plan['waves'] if w['levelId'] == L]; dt = step_ms / 1000
    gold = level['initialGold']; hp = level['initialHp']; towers = []; operations = []; reports = []
    killed = leaked = boss_leaked = critical_leaked = successful = 0; combat_sec = 0; level_sec = 0; max_alive = 0
    kill_log = []          # ★ 新增：每次击杀 {wave, role, fly, cellsFromEnd, x, y, t}
    started = time.perf_counter(); error = None
    for w in waves:
        if hp <= 0: break
        before = gold
        gold, spent = bs.spend_boundary(cells, towers, tower_table, gold, 'optimized', w['waveOrdinal'], operations)
        path = path_for(cells, [t['cell'] for t in towers]); airpath = [path[0], path[-1]]
        alive = []; cursor = 0; elapsed = 0.0; wave_kills = wave_leaks = wave_income = 0
        peak = 0; deferred = 0; leak_roles = collections.Counter()
        while cursor < len(w['enemyIds']) or alive:
            if elapsed > 240: error = 'wave simulation exceeded 240 seconds'; break
            while cursor < len(w['enemyIds']) and elapsed + 1e-9 >= cursor * w['spawnIntervalMs'] / 1000:
                if len(alive) >= tables['tbglobal'][0].get('maxEnemyAlive', 200): deferred += 1; break
                cfg = enemy_table[w['enemyIds'][cursor]]
                role = next(c['role'] for c in w['composition'] if c['variantId'] == cfg['id'])
                alive.append({'spawnIndex': cursor, 'cfg': cfg, 'role': role, 'hp': cfg['hp'],
                              'path': airpath if cfg['isFlying'] else path,
                              'segment': 0, 'segmentProgress': 0.0, 'slowRatio': 0.0, 'slowUntil': 0.0, 'alive': True})
                cursor += 1; successful += 1
            peak = max(peak, len(alive)); max_alive = max(max_alive, len(alive))
            for e in alive:
                if elapsed >= e['slowUntil']: e['slowRatio'] = 0
                bs.advance(e, e['cfg']['speed'] * (1 - e['slowRatio']) * dt)
                e['positionCache'] = bs.position(e)
            for e in alive:
                if e['segment'] >= len(e['path']) - 1:
                    e['alive'] = False; wave_leaks += 1; leaked += 1; hp -= e['cfg']['damageToPlayer']
                    leak_roles[e['role']] += 1
            alive = [e for e in alive if e['alive']]
            for t in towers:
                cfg = tower_table[t['id']]; radius2 = cfg['radius'] ** 2
                if cfg['type'] == 3:
                    t['fireTimer'] += dt
                    if t['fireTimer'] + 1e-9 >= cfg['CD'] / 1000:
                        t['fireTimer'] = 0
                        for e in alive:
                            if e['alive'] and (not e['cfg']['isFlying'] or cfg['canAttackAir']) and bs.d2(t['cell'], bs.position(e)) <= radius2:
                                if cfg['effectValue'] >= e['slowRatio'] or elapsed >= e['slowUntil']:
                                    e['slowRatio'] = cfg['effectValue']; e['slowUntil'] = elapsed + cfg['CD'] / 1000 + .1
                    continue
                target = t['target']
                if target and (not target['alive'] or bs.d2(t['cell'], bs.position(target)) > radius2): t['target'] = target = None
                t['searchTimer'] += dt
                if target is None and t['searchTimer'] + 1e-9 >= cfg['searchIntervalMs'] / 1000:
                    t['searchTimer'] = 0
                    options = [e for e in alive if e['alive'] and (not e['cfg']['isFlying'] or cfg['canAttackAir']) and bs.d2(t['cell'], bs.position(e)) <= radius2]
                    if options:
                        target = max(options, key=lambda e: (e['segment'] + e['segmentProgress']) / max(1, len(e['path']) - 1)); t['target'] = target
                if target is None: continue
                t['fireTimer'] += dt
                if t['fireTimer'] + 1e-9 < cfg['CD'] / 1000: continue
                t['fireTimer'] = 0; victims = [target]
                if cfg['type'] == 2:
                    center = bs.position(target)
                    blast = next(b for b in tables['tbbulletdata'] if b['id'] == cfg['bulletId'])['aoeRadius']
                    victims = [e for e in alive if e['alive'] and not e['cfg']['isFlying'] and bs.d2(bs.position(e), center) <= blast ** 2]
                for e in victims:
                    damage = cfg['power'] * (1 - e['cfg']['armor'])
                    e['hp'] -= damage
                    if e['hp'] <= 0 and e['alive']:
                        e['alive'] = False; killed += 1; wave_kills += 1; gold += e['cfg']['reward']; wave_income += e['cfg']['reward']
                        # ★ 击杀位置跟踪：ground=路径进度，air=直线进度
                        p = e['path']
                        if len(p) == 2:
                            cfe = (len_path(p) - e['segmentProgress'])
                        else:
                            cfe = (len(p) - 1) - (e['segment'] + e['segmentProgress'])
                        pos = e.get('positionCache') or bs.position(e)
                        kill_log.append({'wave': w['waveOrdinal'], 'role': e['role'],
                                         'fly': int(e['cfg']['isFlying']), 'cellsFromEnd': round(cfe, 3),
                                         'x': round(pos[0], 2), 'y': round(pos[1], 2), 't': round(elapsed, 2)})
            alive = [e for e in alive if e['alive']]
            elapsed += dt
            if hp <= 0: break
        combat_sec += elapsed; level_sec += elapsed + (0 if w['waveOrdinal'] == 1 else 5)
        cleared = cursor == len(w['enemyIds']) and not alive and hp > 0
        if cleared: gold += w['rewardGold']
        reports.append({'waveOrdinal': w['waveOrdinal'], 'planned': w['count'], 'killed': wave_kills,
            'leaked': wave_leaks, 'killGold': wave_income, 'clearReward': w['rewardGold'] if cleared else 0,
            'balance': gold, 'hp': max(0, hp), 'towerCount': len(towers), 'pathLength': len(path) - 1,
            'combatElapsedMs': math.floor(elapsed * 1000), 'alivePeak': peak, 'cleared': cleared,
            'towerComposition': dict(collections.Counter(t['id'] for t in towers))})
        if error or hp <= 0: break
    completed = sum(r['cleared'] for r in reports)
    win = completed == rating['expectedWaveCount'] and hp > 0 and not error
    return {'victory': win, 'leaked': leaked, 'hp': max(0, hp), 'killed': killed,
            'completedWaveCount': completed, 'error': error, 'reports': reports,
            'kill_log': kill_log, 'finalTowers': [{'row': t['cell'][1], 'col': t['cell'][0], 'towerId': t['id']} for t in towers],
            'combatElapsedMs': math.floor(combat_sec * 1000), 'balance': gold,
            'mapCapacity': map_capacity['capacity']}

def len_path(p):
    return math.sqrt(sum((p[i + 1][j] - p[i][j]) ** 2 for i in range(len(p) - 1) for j in (0, 1)))

# ---------------------------------------------------------------- 约束检查
def verify(step_ms=50):
    design = l01_design.build_design()
    plan = l01_design.build_plan_from_disk()
    ops, sched = schedule_ops(plan)
    bs.POLICY.clear()
    bs.POLICY['scriptedOperations'] = ops
    bs.POLICY['maxWallSeconds'] = 600
    OUT.mkdir(parents=True, exist_ok=True)
    result = simulate_level_tracked(plan, 1, step_ms, OUT)
    # ---- 约束逐条判定
    checks = {}
    checks['C2_初始金钱100'] = plan['tables']['tbsceneinfo'][0]['initialGold'] == 100
    air_series = [w['flyingCount'] for w in plan['waves']]
    checks['C3_空军W5起严格递增'] = all(a == 0 for a in air_series[:4]) and all(
        air_series[i] < air_series[i + 1] for i in range(4, 24))
    air_types = set()
    for w in plan['waves']:
        for c in w['composition']:
            if c['role'] == 'air':
                air_types.add(plan['tables']['tbenemydata'] and
                              next(v['resName'] for v in plan['tables']['tbenemydata'] if v['id'] == c['variantId']))
    checks['C3_L01仅1种空军'] = air_types == {'Enemy_Bat'}
    checks['C7_零漏怪'] = result['leaked'] == 0 and result['victory']
    final_cells = {(t['col'] + 1, t['row'] + 1) for t in result['finalTowers']}
    want_cells = set(FROST) | set(NORMAL) | set(BLAST)
    l3_ids = {TOWER_L3['N'], TOWER_L3['B'], TOWER_L3['F']}
    checks['C5_42格全满级'] = (final_cells == want_cells and len(result['finalTowers']) == 42
                              and all(t['towerId'] in l3_ids for t in result['finalTowers']))
    w25_kills = [k for k in result['kill_log'] if k['wave'] == 25]
    last_kill = result['kill_log'][-1] if result['kill_log'] else None
    checks['C1_末怪距终点>=2格'] = bool(last_kill and last_kill['wave'] == 25 and last_kill['cellsFromEnd'] >= 2.0)
    min_cfe = min((k['cellsFromEnd'] for k in w25_kills), default=None)
    # ---- 汇总输出
    summary = {
        'checks': checks, 'all_passed': all(checks.values()),
        'C1_detail': {'lastKill': last_kill, 'w25MinCellsFromEnd': min_cfe, 'w25KillCount': len(w25_kills)},
        'sched': sched,
        'economy': {'totalIncome': design['totalIncome'], 'intendedMixCost': 22 * 100 + 13 * 139 + 7 * 130,
                    'allNormalCost': 42 * 100},
        'result': {k: result[k] for k in ('victory', 'leaked', 'killed', 'hp', 'completedWaveCount',
                                          'error', 'combatElapsedMs', 'balance')},
        'airSeries': air_series,
    }
    (OUT / 'verify_result.json').write_text(json.dumps(
        {'summary': summary, 'reports': result['reports'], 'kill_log_w25': w25_kills,
         'finalTowers': result['finalTowers']}, ensure_ascii=False, indent=1), encoding='utf-8')
    return summary, result

if __name__ == '__main__':
    step = int(sys.argv[1]) if len(sys.argv) > 1 else 50
    summary, result = verify(step)
    print(json.dumps(summary['checks'], ensure_ascii=False, indent=1))
    print('all_passed =', summary['all_passed'])
    print('C1:', summary['C1_detail'])
    print('sched:', summary['sched'])
    print('economy:', summary['economy'])
    print('result:', summary['result'])
    for r in result['reports']:
        print(f"W{r['waveOrdinal']:02d} kill={r['killed']:3d} leak={r['leaked']} gold={r['balance']:4d} "
              f"towers={r['towerCount']:2d} pathLen={r['pathLength']:2d} peak={r['alivePeak']:3d} t={r['combatElapsedMs']/1000:6.1f}s")

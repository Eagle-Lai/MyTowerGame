# -*- coding: utf-8 -*-
"""重封 runtimeContentHash —— 严格复刻 Configs.CalculateRuntimeContentHash 的 C# 算法。

步骤：读 Assets/ConfigJson 下 8 个 JSON 原文 + tblevelrating.json 的 30 行 17 字段 →
sha256 小写十六进制 → 写入全部 30 关 Levels/Lxx/LevelRating.xlsx 的 runtimeContentHash 列。
"""
import hashlib, json, pathlib, sys
import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
CFG = ROOT / 'Assets/ConfigJson'
DATA = ROOT / 'Luban/Config/Datas'

def calculate():
    keys = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
            'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal']
    text = []
    for key in keys:
        raw = (CFG / f'{key}.json').read_text(encoding='utf-8-sig')
        text.append(key + '\n' + raw.replace('\r\n', '\n').replace('\r', '\n') + '\n')
    rows = json.loads((CFG / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows = sorted(rows, key=lambda r: r['id'])
    for r in rows:
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  ','.join(str(x) for x in (r.get('bossVariantIds') or [])),
                  ','.join(str(x) for x in (r.get('criticalVariantIds') or [])),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(f) for f in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()

def reseal(hash_value, dry_run=False):
    written = []
    for L in range(1, 31):
        p = DATA / 'Levels' / f'L{L:02d}' / 'LevelRating.xlsx'
        wb = openpyxl.load_workbook(p)
        ws = wb.worksheets[0]
        heads = {c.value: c.column for c in ws[1] if c.value}
        old = None
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value == L:
                old = ws.cell(r, heads['runtimeContentHash']).value
                ws.cell(r, heads['runtimeContentHash'], hash_value)
                break
        if not dry_run:
            wb.save(p)
        wb.close()
        written.append((L, str(old)[:12] if old else None, hash_value[:12]))
    return written

if __name__ == '__main__':
    h = calculate()
    print('新 runtimeContentHash =', h)
    dry = '--dry' in sys.argv
    rows = reseal(h, dry)
    print(f'{"dry-run" if dry else "已写入"} {len(rows)} 关：')
    for L, old, new in rows:
        mark = '' if old == new else '  <- 变更'
        print(f'  L{L:02d}: {old} -> {new}{mark}')

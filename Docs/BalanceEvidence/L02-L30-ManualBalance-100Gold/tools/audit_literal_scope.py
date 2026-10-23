"""Check source changes against explicit cell decisions and preserve workbook structure."""
from formal_tools import *
import openpyxl,collections,sys
tag=sys.argv[1];assert tag in ['r2','r3','r4','final']
names=['design-L02-r1.cells.json','design-L03-L06-r1.cells.json','design-L07-L12-r1.cells.json','design-L13-L18-r1.cells.json','design-L19-L24-r1.cells.json','design-L25-L30-r1.cells.json','rating-identity-r1.cells.json','seal-r1.cells.json','seal-campaign-r1.cells.json','design-air-refinement-r2.cells.json','rating-identity-r2.cells.json','seal-r2.cells.json']
if tag in['r3','r4','final']:
    names += [name for name in ['design-ground-refinement-r3.cells.json','rating-identity-r3.cells.json','seal-r3.cells.json']if(REPORT/name).exists()]
if tag in['r4','final']:
    names += [name for name in ['design-L10-refinement-r4.cells.json','rating-identity-r4.cells.json','seal-r4.cells.json']if(REPORT/name).exists()]
if tag=='final':names+=['rating-calibration-final.cells.json','seal-final.cells.json']
expected={};bybook=collections.defaultdict(set)
for name in names:
    for e in json.loads((REPORT/name).read_text(encoding='utf-8')):
        key=(e['path'],e['sheet'],e['cell']);expected[key]=e['after'];bybook[e['path']].add((e['sheet'],e['cell']))
original=json.loads((REPORT/'audit/protected-files-before.json').read_text(encoding='utf-8'));errors=[];counts=collections.Counter();changed=[]
for path,allowed in bybook.items():
    p=ROOT/path;backup=REPORT/'audit/source-backups'/p.relative_to(ROOT/'Luban/Config/Datas')
    assert backup.exists(),backup
    assert hashlib.sha256(backup.read_bytes()).hexdigest()==original[path],f'Original backup identity mismatch: {path}'
    old=openpyxl.load_workbook(backup);new=openpyxl.load_workbook(p)
    if old.sheetnames!=new.sheetnames:errors.append(dict(path=path,problem='sheet structure'))
    for s in old.sheetnames:
        a,b=old[s],new[s]
        for key in ['max_row','max_column','freeze_panes']:
            if getattr(a,key)!=getattr(b,key):errors.append(dict(path=path,sheet=s,problem=key))
        if str(a.merged_cells)!=str(b.merged_cells):errors.append(dict(path=path,sheet=s,problem='merged ranges'))
        if {k:v.width for k,v in a.column_dimensions.items()}!={k:v.width for k,v in b.column_dimensions.items()}:errors.append(dict(path=path,sheet=s,problem='column widths'))
        if {k:v.height for k,v in a.row_dimensions.items()}!={k:v.height for k,v in b.row_dimensions.items()}:errors.append(dict(path=path,sheet=s,problem='row heights'))
        for rows in a.iter_rows():
            for ca in rows:
                cb=b[ca.coordinate];counts['checkedCells']+=1
                if ca._style!=cb._style:errors.append(dict(path=path,sheet=s,cell=ca.coordinate,problem='style changed'))
                if ca.value!=cb.value:
                    counts['changedCells']+=1
                    if(s,ca.coordinate)not in allowed:errors.append(dict(path=path,sheet=s,cell=ca.coordinate,problem='outside explicit decisions'))
                    changed.append(dict(path=path,sheet=s,cell=ca.coordinate,before=ca.value,after=cb.value))
        for sheet,cell in allowed:
            if sheet==s and b[cell].value!=expected[path,sheet,cell]:errors.append(dict(path=path,sheet=s,cell=cell,problem='not the last decided literal',actual=b[cell].value,expected=expected[path,sheet,cell]))
    counts['workbooks']+=1
record=dict(tag=tag,originalBackupIdentitiesMatch=True,explicitManifests=names,counts=dict(counts),l01OnlyRuntimeSeal=all(x['path']!='Luban/Config/Datas/Levels/L01/LevelRating.xlsx'or(x['sheet']=='data'and x['cell']=='T4')for x in changed),errors=errors,actualChanges=changed)
(REPORT/f'audit/literal-scope-{tag}.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print(dict(counts),len(errors),'errors');assert not errors and record['l01OnlyRuntimeSeal']

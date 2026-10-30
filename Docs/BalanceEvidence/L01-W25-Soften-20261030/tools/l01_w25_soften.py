# -*- coding: utf-8 -*-
"""L01 W25（终局波）软着陆：地面/空中血量 −20%，出怪间隔 +10%（2026-10-30 用户需求）。

需求原文：
  「减少最后一个波次的地面敌人的血量 20%，空中敌人的血量 20%，出怪间隙增加 10% 的时间。」

落地口径（全部在当前表值上做增量）：
- W25 的 6 个变体（5000240..5000245，5 地面 + 1 空中）hp 各 × 0.80；
  数量、护甲、速度、击杀金一律不动。
- W25 的 enemyInterval 215ms → **236.5ms**（×1.10，字段是 float，保留一位小数）。
- LevelRating 只更新 planHash / balanceRevision（敌数不变，仍 695）。

用法：
  python l01_w25_soften.py values | reseal | verify
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
L01 = DATAS / 'Levels/L01'
JSON_DIR = ROOT / 'Assets/ConfigJson'
TAG = 'w25soften-20261030'
REVISION = 'w25-soften-20261030-v1'
HP_SCALE = 0.80
INTERVAL_SCALE = 1.10
W25_LIST_ID = 405


def load_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    rows = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            rows.append((r, {k: ws.cell(r, c).value for k, c in heads.items()}))
    return wb, ws, heads, rows


def ehp(seq, variants):
    return sum(float(variants[v]['hp']) / (1 - float(variants[v]['armor'] or 0)) for v in seq)


def cmd_values():
    _, _, _, list_rows = load_rows(L01 / 'EnemyList.xlsx')
    lists = {int(v['id']): v for _, v in list_rows}
    _, _, _, var_rows = load_rows(L01 / 'EnemyData.xlsx')
    variants = {int(v['id']): v for _, v in var_rows}
    seq = [int(x) for x in str(lists[W25_LIST_ID]['EnemyIndexs']).split(',')]
    ids = sorted(set(seq))
    old_interval = float(lists[W25_LIST_ID]['enemyInterval'])
    if abs(old_interval - 215.0) > 1e-6:
        raise RuntimeError('基线不符：W25 interval = %s，应为 215' % old_interval)
    if any(TAG in str(variants[i]['desc'] or '') for i in ids):
        raise RuntimeError('已应用 W25 软着陆，拒绝重复执行')

    old_ehp = ehp(seq, variants)
    old_air = ehp([v for v in seq if int(variants[v]['isFlying'] or 0) == 1], variants)

    # 1) 血量 ×0.8（地面 + 空中一起）
    wb, ws, heads, rows = load_rows(L01 / 'EnemyData.xlsx')
    for r, v in rows:
        vid = int(v['id'])
        if vid not in ids:
            continue
        old_hp = float(v['hp'])
        ws.cell(r, heads['hp'], round(old_hp * HP_SCALE, 6))
        d = str(v['desc'] or '').split(' | ' + TAG)[0]
        ws.cell(r, heads['desc'], d + ' | %s: hp x%.2f' % (TAG, HP_SCALE))
        print('  %d %-8s %s hp %.3f -> %.3f' % (vid, v['name'], '空中' if int(v['isFlying'] or 0) else '地面',
                                                old_hp, old_hp * HP_SCALE))
    wb.save(L01 / 'EnemyData.xlsx')
    wb.close()

    # 2) 出怪间隔 ×1.1
    _, _, _, var_rows2 = load_rows(L01 / 'EnemyData.xlsx')
    variants2 = {int(v['id']): v for _, v in var_rows2}
    new_ehp = ehp(seq, variants2)
    new_air = ehp([v for v in seq if int(variants2[v]['isFlying'] or 0) == 1], variants2)
    new_interval = round(old_interval * INTERVAL_SCALE, 2)

    wb, ws, heads, rows = load_rows(L01 / 'EnemyList.xlsx')
    for r, v in rows:
        if int(v['id']) != W25_LIST_ID:
            continue
        ws.cell(r, heads['enemyInterval'], new_interval)
        parts = str(v['desc'] or '').split(' | ')
        parts[0] = 'L1 W25/25; N=%d; EHP=%.6f' % (len(seq), new_ehp)
        parts = [p for p in parts if TAG not in p]
        ws.cell(r, heads['desc'], ' | '.join(parts) + ' | %s: hp x%.2f, interval x%.2f' % (TAG, HP_SCALE, INTERVAL_SCALE))
    wb.save(L01 / 'EnemyList.xlsx')
    wb.close()

    old_span = len(seq) * old_interval / 1000.0
    new_span = len(seq) * new_interval / 1000.0
    print('\nW25 敌数 %d（地面 %d / 空中 %d）' % (len(seq), len(seq) - sum(1 for v in seq if int(variants2[v]['isFlying'] or 0)), sum(1 for v in seq if int(variants2[v]['isFlying'] or 0))))
    print('总 EHP %.0f -> %.0f（%+.1f%%，地面+空中各 -20%%）' % (old_ehp, new_ehp, (new_ehp / old_ehp - 1) * 100))
    print('空军 EHP %.0f -> %.0f（%+.1f%%）' % (old_air, new_air, (new_air / old_air - 1) * 100))
    print('出怪间隔 %.1fms -> %.1fms（%+.1f%%）；出怪总时长 %.1fs -> %.1fs（%+.1f%%）'
          % (old_interval, new_interval, (new_interval / old_interval - 1) * 100, old_span, new_span, (new_span / old_span - 1) * 100))
    print('压力 EHP/秒 %.0f -> %.0f（%+.1f%%）' % (old_ehp / old_span, new_ehp / new_span, ((new_ehp / new_span) / (old_ehp / old_span) - 1) * 100))

    # 3) LevelRating
    spec = json.dumps({'mode': 'w25-soften', 'hp': HP_SCALE, 'interval': INTERVAL_SCALE}, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    wb, ws, heads, rows = load_rows(L01 / 'LevelRating.xlsx')
    for r, v in rows:
        ws.cell(r, heads['planHash'], plan_hash)
        ws.cell(r, heads['balanceRevision'], REVISION)
    wb.save(L01 / 'LevelRating.xlsx')
    wb.close()
    print('LevelRating: planHash=%s…, revision=%s（敌数不变）' % (plan_hash[:12], REVISION))


def compute_hash():
    keys = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
            'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal']
    text = []
    for k in keys:
        p = JSON_DIR / (k + '.json')
        s = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        text.append('%s\n%s\n' % (k, s))
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows.sort(key=lambda r: r['id'])
    for r in rows:
        def joined(v):
            return ','.join(str(x) for x in v) if isinstance(v, list) else ''
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  joined(r.get('bossVariantIds')), joined(r.get('criticalVariantIds')),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(x) for x in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()


def cmd_reseal():
    h = compute_hash()
    print('按导出 JSON 计算 runtimeContentHash = %s' % h)
    for i in range(1, 31):
        p = DATAS / ('Levels/L%02d/LevelRating.xlsx' % i)
        wb, ws, heads, rows = load_rows(p)
        for r, v in rows:
            ws.cell(r, heads['runtimeContentHash'], h)
        wb.save(p)
        wb.close()
    print('已写入 30 关 LevelRating.xlsx（runtimeContentHash）')


def cmd_verify():
    h = compute_hash()
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    stored = sorted({r['runtimeContentHash'] for r in rows})
    print('重算哈希: %s 一致=%s' % (h, stored == [h]))
    enemy = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemydata.json').read_text(encoding='utf-8-sig'))}
    lists = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemylist.json').read_text(encoding='utf-8-sig'))}
    w = lists[W25_LIST_ID]
    seq = w['EnemyIndexs']
    air = [v for v in seq if enemy[v]['isFlying'] == 1]
    ground = [v for v in seq if enemy[v]['isFlying'] == 0]
    e = lambda arr: sum(float(enemy[v]['hp']) / (1 - float(enemy[v]['armor'] or 0)) for v in arr)
    print('运行时 JSON：W25 敌数 %d（地面 %d / 空中 %d）间隔 %.2fms' % (len(seq), len(ground), len(air), w['enemyInterval']))
    print('  地面 EHP %.0f｜空中 EHP %.0f｜合计 %.0f' % (e(ground), e(air), e(seq)))


if __name__ == '__main__':
    mode = sys.argv[1] if len(sys.argv) > 1 else 'verify'
    {'values': cmd_values, 'reseal': cmd_reseal, 'verify': cmd_verify}[mode]()

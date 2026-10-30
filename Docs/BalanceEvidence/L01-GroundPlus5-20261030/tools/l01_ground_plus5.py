# -*- coding: utf-8 -*-
"""L01：W5 起地面怪数量 +5%（2026-10-30 用户需求）。

需求原文：
  「第一关给波次，从第五波开始，地面怪物的数量都增加 5%。」

落地口径：
- 只加**地面**（isFlying=0），空军数量与序列完全不动（保持"空军逐波递增"的设计）。
- 每波新增数量 = round(该波地面数 × 0.05)；新增个体按**该波现有地面变体占比**分配
  （最大余数法，保证成分不变、总数精确）。
- 新增个体按均匀间隔插入波次序列（不是全堆在末尾），出怪节奏尽量均匀。
- LevelRating.expectedEnemyCount 同步增加；planHash / balanceRevision 更新。

用法：
  python l01_ground_plus5.py values | reseal | verify
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
L01 = DATAS / 'Levels/L01'
JSON_DIR = ROOT / 'Assets/ConfigJson'
TAG = 'ground5-20261030'
REVISION = 'ground-plus5-20261030-v1'
RATIO = 0.05
FROM_WAVE = 5


def load_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    rows = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            rows.append((r, {k: ws.cell(r, c).value for k, c in heads.items()}))
    return wb, ws, heads, rows


def ehp(seq, variants):
    return sum(float(variants[v]['hp']) / (1 - float(variants[v]['armor'] or 0)) for v in seq)


def allocate(counts, add):
    """按占比分配 add 个名额（最大余数法），返回 {vid: n}，总数精确等于 add。"""
    total = sum(counts.values())
    if total <= 0 or add <= 0:
        return {}
    exact = {v: c * add / total for v, c in counts.items()}
    base = {v: int(x) for v, x in exact.items()}
    rest = add - sum(base.values())
    for v in sorted(exact, key=lambda k: (-(exact[k] - base[k]), k))[:rest]:
        base[v] += 1
    return {v: n for v, n in base.items() if n > 0}


def splice(seq, extras):
    """把 extras 均匀插进 seq（保持相对顺序），避免全部堆在末尾。"""
    if not extras:
        return list(seq)
    out = []
    step = len(seq) / float(len(extras) + 1)
    next_insert = step
    m = 0
    for i, v in enumerate(seq):
        out.append(v)
        if m < len(extras) and (i + 1) >= next_insert:
            out.append(extras[m])
            m += 1
            next_insert = step * (m + 1)
    out.extend(extras[m:])
    return out


def cmd_values():
    _, _, _, scene = load_rows(L01 / 'SceneInfo.xlsx')
    order = [int(x) for x in str(scene[0][1]['RoundList']).split(',')]
    _, _, _, round_rows = load_rows(L01 / 'RoundData.xlsx')
    rounds = {int(v['id']): v for _, v in round_rows}
    _, _, _, list_rows = load_rows(L01 / 'EnemyList.xlsx')
    lists = {int(v['id']): v for _, v in list_rows}
    _, _, _, var_rows = load_rows(L01 / 'EnemyData.xlsx')
    variants = {int(v['id']): v for _, v in var_rows}
    if any(TAG in str(v.get('desc') or '') for v in lists.values()):
        raise RuntimeError('已应用地面 +5%，拒绝重复执行')

    lid = {w: int(str(rounds[order[w - 1]]['EnemyIndexs']).split(',')[0]) for w in range(1, 26)}
    wb, ws, heads, rows = load_rows(L01 / 'EnemyList.xlsx')
    by_id = {int(v['id']): (r, v) for r, v in rows}
    total_add = 0
    print('波次  地面(旧→新)  新增  总敌数(旧→新)  EHP(旧→新)')
    for w in range(FROM_WAVE, 26):
        r, v = by_id[lid[w]]
        seq = [int(x) for x in str(v['EnemyIndexs']).split(',')]
        ground = [x for x in seq if int(variants[x]['isFlying'] or 0) == 0]
        counts = {}
        for x in ground:
            counts[x] = counts.get(x, 0) + 1
        add = int(round(len(ground) * RATIO))
        plan = allocate(counts, add)
        extras = []
        for vid in sorted(plan):
            extras.extend([vid] * plan[vid])
        new_seq = splice(seq, extras)
        new_ground = len([x for x in new_seq if int(variants[x]['isFlying'] or 0) == 0])
        assert new_ground == len(ground) + add, '地面数量不匹配'
        ws.cell(r, heads['EnemyIndexs'], ','.join(str(x) for x in new_seq))
        parts = str(v.get('desc') or '').split(' | ')
        parts[0] = 'L1 W%d/25; N=%d; EHP=%.6f' % (w, len(new_seq), ehp(new_seq, variants))
        parts = [p for p in parts if TAG not in p]
        ws.cell(r, heads['desc'], ' | '.join(parts) + ' | %s: ground +%d (%.0f%%)' % (TAG, add, RATIO * 100))
        total_add += add
        if w <= 6 or w in (10, 15, 20, 21, 24, 25):
            print('W%02d   %4d→%-4d  +%-3d  %4d→%-4d  %.0f→%.0f'
                  % (w, len(ground), new_ground, add, len(seq), len(new_seq), ehp(seq, variants), ehp(new_seq, variants)))
    wb.save(L01 / 'EnemyList.xlsx')
    wb.close()
    print('\nW5-W25 地面新增合计 %d 只' % total_add)

    # LevelRating：敌数同步 + planHash
    spec = json.dumps({'mode': 'ground-plus5', 'ratio': RATIO, 'fromWave': FROM_WAVE, 'added': total_add}, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    wb, ws, heads, rows = load_rows(L01 / 'LevelRating.xlsx')
    for r, v in rows:
        old = int(v['expectedEnemyCount'] or 0)
        ws.cell(r, heads['expectedEnemyCount'], old + total_add)
        ws.cell(r, heads['planHash'], plan_hash)
        ws.cell(r, heads['balanceRevision'], REVISION)
        print('LevelRating: expectedEnemyCount %d -> %d, planHash=%s…, revision=%s'
              % (old, old + total_add, plan_hash[:12], REVISION))
    wb.save(L01 / 'LevelRating.xlsx')
    wb.close()

    # 复核：W20 的相对增幅（用户上一轮把 W20 定在 +30%，地面 +5% 会把它顶上去）
    _, _, _, list_rows2 = load_rows(L01 / 'EnemyList.xlsx')
    lists2 = {int(v['id']): v for _, v in list_rows2}
    e = {}
    for w in (19, 20, 21, 24, 25):
        seq = [int(x) for x in str(lists2[lid[w]]['EnemyIndexs']).split(',')]
        e[w] = ehp(seq, variants)
    print('\n复查相对增幅：W19→W20 %+.1f%%（改动前 +30.0%%）｜W20→W21 %+.1f%%｜W24→W25 %+.1f%%'
          % ((e[20] / e[19] - 1) * 100, (e[21] / e[20] - 1) * 100, (e[25] / e[24] - 1) * 100))


def compute_hash():
    keys = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
            'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal']
    text = []
    for k in keys:
        p = JSON_DIR / (k + '.json')
        s = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        text.append('%s\n%s\n' % (k, s))
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows.sort(key=lambda r: r['id'])
    for r in rows:
        def joined(v):
            return ','.join(str(x) for x in v) if isinstance(v, list) else ''
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  joined(r.get('bossVariantIds')), joined(r.get('criticalVariantIds')),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(x) for x in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()


def cmd_reseal():
    h = compute_hash()
    print('按导出 JSON 计算 runtimeContentHash = %s' % h)
    for i in range(1, 31):
        p = DATAS / ('Levels/L%02d/LevelRating.xlsx' % i)
        wb, ws, heads, rows = load_rows(p)
        for r, v in rows:
            ws.cell(r, heads['runtimeContentHash'], h)
        wb.save(p)
        wb.close()
    print('已写入 30 关 LevelRating.xlsx（runtimeContentHash）')


def cmd_verify():
    h = compute_hash()
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    stored = sorted({r['runtimeContentHash'] for r in rows})
    print('重算哈希: %s 一致=%s' % (h, stored == [h]))
    enemy = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemydata.json').read_text(encoding='utf-8-sig'))}
    lists = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemylist.json').read_text(encoding='utf-8-sig'))}
    rounds = {x['id']: x for x in json.loads((JSON_DIR / 'tbrounddata.json').read_text(encoding='utf-8-sig'))}
    lv = [x for x in json.loads((JSON_DIR / 'tbsceneinfo.json').read_text(encoding='utf-8-sig')) if x['id'] == 1][0]
    total = ground = air = 0
    for w, rid in enumerate(lv['RoundList'], 1):
        seq = lists[rounds[rid]['EnemyIndexs'][0]]['EnemyIndexs']
        g = [v for v in seq if enemy[v]['isFlying'] == 0]
        total += len(seq); ground += len(g); air += len(seq) - len(g)
        if w in (4, 5, 20, 21, 25):
            print('W%02d 敌数 %d（地面 %d / 空中 %d）' % (w, len(seq), len(g), len(seq) - len(g)))
    print('全关：敌数 %d（地面 %d / 空中 %d）｜rating.expectedEnemyCount=%d'
          % (total, ground, air, rows[0]['expectedEnemyCount']))


if __name__ == '__main__':
    mode = sys.argv[1] if len(sys.argv) > 1 else 'verify'
    {'values': cmd_values, 'reseal': cmd_reseal, 'verify': cmd_verify}[mode]()

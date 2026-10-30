# -*- coding: utf-8 -*-
"""L01：W5 起波次奖励 +5/+10、W21-W25 空军血量 −5% 且每波少 5 只（2026-10-29 用户需求）。

需求原文：
  「从第五波次后，每次通关奖励增加 5-10 个金币。希望第 21 个波次，空中怪物的血量减少 5%，
    往后的关卡也减少 5%，并且减少 5 个怪物。」

落地口径：
1) 波次通关奖励（RoundData.rewardGold）：W5 起（含）—— 压力波 W5/W10/W15/W20/W25 **+10**，
   其余 W6-W24 的普通波 **+5**；W1-W4 不动。合计 +130 金。
2) 空军血量：W21-W25 的 Bat 变体 hp × 0.95（−5%）。
3) 出怪数量：W21-W25 每波**少 5 只**，从该波序列尾部剔除空军（蝙蝠），保证地面节奏不变。
   → 副作用：空军逐波递增的规则在 W20(68) → W21(67) 处被打破 1 只（如需修复可再 trim W20 或改剔地面）。
4) LevelRating.expectedEnemyCount 同步 −25（9820 → 9795）；planHash / balanceRevision 更新。

用法：
  python l01_wavegold_air.py values | reseal | verify
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
L01 = DATAS / 'Levels/L01'
JSON_DIR = ROOT / 'Assets/ConfigJson'
TAG = 'wavegold-air-20261029'
REVISION = 'wavegold-air-20261029-v1'
PRESSURE_WAVES = (5, 10, 15, 20, 25)
BONUS_PRESSURE = 10
BONUS_NORMAL = 5
AIR_HP_SCALE = 0.95
AIR_TRIM = 5
AIR_FROM_WAVE = 21


def load_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    rows = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            rows.append((r, {k: ws.cell(r, c).value for k, c in heads.items()}))
    return wb, ws, heads, rows


def read_level():
    _, _, _, scene = load_rows(L01 / 'SceneInfo.xlsx')
    order = [int(x) for x in str(scene[0][1]['RoundList']).split(',')]
    _, _, _, rounds_rows = load_rows(L01 / 'RoundData.xlsx')
    rounds = {int(v['id']): v for _, v in rounds_rows}
    _, _, _, list_rows = load_rows(L01 / 'EnemyList.xlsx')
    lists = {int(v['id']): v for _, v in list_rows}
    _, _, _, var_rows = load_rows(L01 / 'EnemyData.xlsx')
    variants = {int(v['id']): v for _, v in var_rows}
    return order, rounds, lists, variants


def ehp_of(seq, variants):
    return sum(float(variants[v]['hp']) / (1 - float(variants[v]['armor'] or 0)) for v in seq)


def cmd_values():
    order, rounds, lists, variants = read_level()
    if any(TAG in str(v.get('desc') or '') for v in lists.values()):
        raise RuntimeError('已应用本批调整，拒绝重复执行')
    lid_of = {w: int(str(rounds[order[w - 1]]['EnemyIndexs']).split(',')[0]) for w in range(1, 26)}

    # 1) 波次奖励
    wb, ws, heads, rows = load_rows(L01 / 'RoundData.xlsx')
    by_id = {int(v['id']): (r, v) for r, v in rows}
    bonus_total = 0
    for w in range(5, 26):
        rid = order[w - 1]
        r, v = by_id[rid]
        add = BONUS_PRESSURE if w in PRESSURE_WAVES else BONUS_NORMAL
        old = int(v['rewardGold'] or 0)
        ws.cell(r, heads['rewardGold'], old + add)
        bonus_total += add
        if w <= 6 or w in PRESSURE_WAVES:
            print('  W%02d 通关奖励 %d -> %d' % (w, old, old + add))
    wb.save(L01 / 'RoundData.xlsx')
    wb.close()
    print('W5-W25 波次奖励合计 +%d 金（压力波 +%d / 普通波 +%d）' % (bonus_total, BONUS_PRESSURE, BONUS_NORMAL))

    # 2) 空军 hp -5%（W21-W25）
    air_ids = {}
    for w in range(AIR_FROM_WAVE, 26):
        seq = [int(x) for x in str(lists[lid_of[w]]['EnemyIndexs']).split(',')]
        air_ids[w] = sorted({v for v in seq if int(variants[v]['isFlying'] or 0) == 1})
    wb, ws, heads, rows = load_rows(L01 / 'EnemyData.xlsx')
    touched = set()
    for r, v in rows:
        vid = int(v['id'])
        if vid in {x for ids in air_ids.values() for x in ids}:
            old = float(v['hp'])
            ws.cell(r, heads['hp'], round(old * AIR_HP_SCALE, 6))
            d = str(v['desc'] or '')
            ws.cell(r, heads['desc'], d.split(' | ' + TAG)[0] + ' | %s: hp x%.2f' % (TAG, AIR_HP_SCALE))
            touched.add(vid)
            print('  %d %-6s hp %.3f -> %.3f' % (vid, v['name'], old, old * AIR_HP_SCALE))
    wb.save(L01 / 'EnemyData.xlsx')
    wb.close()

    # 3) W21-W25 每波剔除 5 只空军（从序列尾部往前删）
    wb, ws, heads, rows = load_rows(L01 / 'EnemyList.xlsx')
    by_list = {int(v['id']): (r, v) for r, v in rows}
    for w in range(AIR_FROM_WAVE, 26):
        r, v = by_list[lid_of[w]]
        seq = [int(x) for x in str(v['EnemyIndexs']).split(',')]
        ground = [x for x in seq if int(variants[x]['isFlying'] or 0) == 0]
        air = [x for x in seq if int(variants[x]['isFlying'] or 0) == 1]
        if len(air) <= AIR_TRIM:
            raise RuntimeError('W%d 空军不足 %d 只，无法剔除' % (w, AIR_TRIM))
        # 重新拼：按原顺序保留地面全部 + 空军前 (n-5) 只
        keep_air = len(air) - AIR_TRIM
        out, seen = [], 0
        for x in seq:
            if int(variants[x]['isFlying'] or 0) == 1:
                seen += 1
                if seen > keep_air:
                    continue
            out.append(x)
        ws.cell(r, heads['EnemyIndexs'], ','.join(str(x) for x in out))
        ws.cell(r, heads['interval'], v['interval'])
        d = str(v['desc'] or '').split(' | ' + TAG)[0]
        ws.cell(r, heads['desc'], '%s | %s: air -%d, air hp x%.2f' % (d, TAG, AIR_TRIM, AIR_HP_SCALE))
        print('  W%02d 敌数 %d -> %d（空军 %d -> %d）' % (w, len(seq), len(out), len(air), keep_air))
    wb.save(L01 / 'EnemyList.xlsx')
    wb.close()

    # 4) LevelRating：敌数 -25、planHash、修订号
    total_trim = AIR_TRIM * (26 - AIR_FROM_WAVE)
    spec = json.dumps({'mode': 'wavegold-air', 'bonus': [BONUS_PRESSURE, BONUS_NORMAL],
                       'pressure': list(PRESSURE_WAVES), 'airHp': AIR_HP_SCALE,
                       'airTrim': [AIR_FROM_WAVE, 25, AIR_TRIM], 'trimTotal': total_trim},
                      ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    wb, ws, heads, rows = load_rows(L01 / 'LevelRating.xlsx')
    for r, v in rows:
        old = int(v['expectedEnemyCount'] or 0)
        ws.cell(r, heads['expectedEnemyCount'], old - total_trim)
        ws.cell(r, heads['planHash'], plan_hash)
        ws.cell(r, heads['balanceRevision'], REVISION)
        print('LevelRating: expectedEnemyCount %d -> %d, planHash=%s…, revision=%s'
              % (old, old - total_trim, plan_hash[:12], REVISION))
    wb.save(L01 / 'LevelRating.xlsx')
    wb.close()

    # 5) 复核输出
    order, rounds, lists, variants = read_level()
    print('\n波次  敌数  空军  空军EHP  总EHP     通关奖励')
    tot_wave = 0
    for w in range(19, 26):
        lid = lid_of[w]
        seq = [int(x) for x in str(lists[lid]['EnemyIndexs']).split(',')]
        air = [v for v in seq if int(variants[v]['isFlying'] or 0) == 1]
        rw = int(rounds[order[w - 1]]['rewardGold'] or 0)
        tot_wave += rw
        print('W%02d  %5d  %3d  %7.0f  %8.0f  %5d' % (w, len(seq), len(air), ehp_of(air, variants), ehp_of(seq, variants), rw))


def compute_hash():
    keys = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
            'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal']
    text = []
    for k in keys:
        p = JSON_DIR / (k + '.json')
        s = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        text.append('%s\n%s\n' % (k, s))
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows.sort(key=lambda r: r['id'])
    for r in rows:
        def joined(v):
            return ','.join(str(x) for x in v) if isinstance(v, list) else ''
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  joined(r.get('bossVariantIds')), joined(r.get('criticalVariantIds')),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(x) for x in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()


def cmd_reseal():
    h = compute_hash()
    print('按导出 JSON 计算 runtimeContentHash = %s' % h)
    for i in range(1, 31):
        p = DATAS / ('Levels/L%02d/LevelRating.xlsx' % i)
        wb, ws, heads, rows = load_rows(p)
        for r, v in rows:
            ws.cell(r, heads['runtimeContentHash'], h)
        wb.save(p)
        wb.close()
    print('已写入 30 关 LevelRating.xlsx（runtimeContentHash）')


def cmd_verify():
    h = compute_hash()
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    stored = sorted({r['runtimeContentHash'] for r in rows})
    print('重算哈希: %s' % h)
    print('表中哈希: %s 一致=%s' % (stored, stored == [h]))
    enemy = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemydata.json').read_text(encoding='utf-8-sig'))}
    lists = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemylist.json').read_text(encoding='utf-8-sig'))}
    rounds = {x['id']: x for x in json.loads((JSON_DIR / 'tbrounddata.json').read_text(encoding='utf-8-sig'))}
    lv = [x for x in json.loads((JSON_DIR / 'tbsceneinfo.json').read_text(encoding='utf-8-sig')) if x['id'] == 1][0]
    kill = 0.0
    for w, rid in enumerate(lv['RoundList'], 1):
        seq = lists[rounds[rid]['EnemyIndexs'][0]]['EnemyIndexs']
        k = sum(float(enemy[v]['reward']) for v in seq)
        kill += k
        if w >= 19:
            air = [v for v in seq if enemy[v]['isFlying'] == 1]
            print('W%02d 敌数 %d 空军 %d 击杀金 %.2f 通关奖 %d' % (w, len(seq), len(air), k, rounds[rid]['rewardGold']))
    wave = sum(rounds[r]['rewardGold'] for r in lv['RoundList'])
    print('敌数合计 %d（rating=%d）| 击杀金 %.2f + 波奖 %d + 初始 100 = 总收入 %.2f'
          % (sum(len(lists[rounds[r]['EnemyIndexs'][0]]['EnemyIndexs']) for r in lv['RoundList']),
             rows[0]['expectedEnemyCount'], kill, wave, 100 + kill + wave))


if __name__ == '__main__':
    mode = sys.argv[1] if len(sys.argv) > 1 else 'verify'
    {'values': cmd_values, 'reseal': cmd_reseal, 'verify': cmd_verify}[mode]()

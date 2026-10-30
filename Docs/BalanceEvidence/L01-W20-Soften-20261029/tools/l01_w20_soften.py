# -*- coding: utf-8 -*-
"""L01 W20 压力波压平（2026-10-29 用户需求）。

需求原文：
  「调整这个波次，全波等效血量 EHP 总增幅控制在 30% 左右，而且出怪间隔改为 −2.5%，
    然后从第二十一波到 25 波次的增幅都维持不变。」

落地口径：
- W20 的 EHP = W19 的 EHP × 1.30（原为 ×1.68，+68.0% → +30.0%）。只按同一系数缩 5 个 W20 变体的 hp，
  组成/护甲/数量/奖励全不动，所以 EHP 精确等比缩到目标。
- W20 出怪间隔 = W19 的 154ms × 0.975 = **150.15ms**（原 147ms，原增幅 −4.5% → −2.5%）。
  `enemyInterval` 的 Luban 类型是 float、代码侧 `SpawnIntervalMs` 也是 float，保留两位小数安全。
- **W21-W25 一律不动**（数值、间隔、奖励都不碰），因此 W20→W21 会从 +2.2% 变成约 +32%（把原来
  一个大台阶摊成两个中等台阶）。若要把整条尾巴等比缩放，另说。
- W20 的敌数与击杀金不变（580 只 / 150.24 金，奖励按角色而非血量），整关总收入仍 3667。

用法：
  python l01_w20_soften.py values   # 改表（Luban 导出前）
  python l01_w20_soften.py reseal   # 按导出 JSON 重算 runtimeContentHash（导出后）
  python l01_w20_soften.py verify   # 只读复核
"""
import hashlib
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
DATAS = ROOT / 'Luban/Config/Datas'
L01 = DATAS / 'Levels/L01'
JSON_DIR = ROOT / 'Assets/ConfigJson'
TAG = 'w20soften-20261029'
REVISION = 'w20-soften-20261029-v1'
EHP_GROWTH = 1.30
INTERVAL_GROWTH = 0.975


def load_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    rows = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            rows.append((r, {k: ws.cell(r, c).value for k, c in heads.items()}))
    return wb, ws, heads, rows


def read_level():
    _, _, _, scene = load_rows(L01 / 'SceneInfo.xlsx')
    order = [int(x) for x in str(scene[0][1]['RoundList']).split(',')]
    _, _, _, rounds_rows = load_rows(L01 / 'RoundData.xlsx')
    rounds = {int(v['id']): v for _, v in rounds_rows}
    _, _, _, list_rows = load_rows(L01 / 'EnemyList.xlsx')
    lists = {int(v['id']): v for _, v in list_rows}
    _, _, _, var_rows = load_rows(L01 / 'EnemyData.xlsx')
    variants = {int(v['id']): v for _, v in var_rows}
    return order, rounds, lists, variants


def seq_of(wave, order, rounds, lists):
    rid = order[wave - 1]
    lid = int(str(rounds[rid]['EnemyIndexs']).split(',')[0])
    return rid, lid, [int(x) for x in str(lists[lid]['EnemyIndexs']).split(',')]


def ehp_of(seq, variants):
    return sum(float(variants[v]['hp']) / (1 - float(variants[v]['armor'] or 0)) for v in seq)


def cmd_values():
    order, rounds, lists, variants = read_level()
    _, _, seq19 = seq_of(19, order, rounds, lists)
    rid20, lid20, seq20 = seq_of(20, order, rounds, lists)
    if len(seq20) != 580:
        raise RuntimeError('基线不符：W20 敌数 %d != 580' % len(seq20))
    if any(TAG in str(variants[v]['desc'] or '') for v in seq20):
        raise RuntimeError('已应用 W20 压平，拒绝重复执行')
    e19, e20 = ehp_of(seq19, variants), ehp_of(seq20, variants)
    target = e19 * EHP_GROWTH
    k = target / e20
    print('W19 EHP=%.3f  W20 EHP=%.3f（%+.1f%%）→ 目标 %.3f（%+.1f%%），缩放系数 %.6f'
          % (e19, e20, (e20 / e19 - 1) * 100, target, EHP_GROWTH * 100 - 100, k))

    # 1) W20 的 5 个变体 hp 等比缩放
    wb, ws, heads, rows = load_rows(L01 / 'EnemyData.xlsx')
    ids20 = set(seq20)
    n_rw = 0
    for r, v in rows:
        vid = int(v['id'])
        if vid not in ids20:
            continue
        old_hp = float(v['hp'])
        new_hp = round(old_hp * k, 6)
        ws.cell(r, heads['hp'], new_hp)
        d = str(ws.cell(r, heads['desc']).value or '')
        ws.cell(r, heads['desc'], d.split(' | ' + TAG)[0] + ' | %s: hp x%.4f' % (TAG, k))
        n_rw += 1
        print('  %d %-8s hp %.3f -> %.3f' % (vid, v['name'], old_hp, new_hp))
    wb.save(L01 / 'EnemyData.xlsx')
    wb.close()

    # 2) W20 出怪间隔：154 x 0.975
    wb, ws, heads, rows = load_rows(L01 / 'EnemyList.xlsx')
    w19 = next(v for _, v in rows if int(v['id']) == lid20 - 1) if lid20 > 1 else None
    base_ms = float(lists[lid20 - 1]['enemyInterval']) if (lid20 - 1) in lists else 154.0
    new_ms = round(base_ms * INTERVAL_GROWTH, 2)
    for r, v in rows:
        if int(v['id']) != lid20:
            continue
        old_ms = float(v['enemyInterval'])
        ws.cell(r, heads['enemyInterval'], new_ms)
        d = str(ws.cell(r, heads['desc']).value or '')
        ws.cell(r, heads['desc'], d.split(' | ' + TAG)[0] + ' | %s: interval %.2fms' % (TAG, new_ms))
        print('W20 出怪间隔 %.0fms -> %.2fms（相对 W19 的 %.0fms：%+.2f%%，原 %+.2f%%）'
              % (old_ms, new_ms, base_ms, (new_ms / base_ms - 1) * 100, (old_ms / base_ms - 1) * 100))
    wb.save(L01 / 'EnemyList.xlsx')
    wb.close()

    # 3) LevelRating：planHash / balanceRevision（敌数、波数、par 不变）
    spec = json.dumps({'mode': 'w20-soften', 'wave': 20, 'ehpGrowth': EHP_GROWTH,
                       'intervalGrowth': INTERVAL_GROWTH, 'k': round(k, 9), 'baseMs': base_ms},
                      ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    wb, ws, heads, rows = load_rows(L01 / 'LevelRating.xlsx')
    for r, v in rows:
        ws.cell(r, heads['planHash'], plan_hash)
        ws.cell(r, heads['balanceRevision'], REVISION)
    wb.save(L01 / 'LevelRating.xlsx')
    wb.close()
    print('L01 LevelRating.xlsx: planHash=%s…, revision=%s' % (plan_hash[:12], REVISION))

    # 4) 复核：重读一遍表，确认只有 W20 变了
    order, rounds, lists, variants = read_level()
    print('\n波次  敌数   EHP       相对上一波   间隔ms')
    prev = None
    for w in range(17, 26):
        rid, lid, seq = seq_of(w, order, rounds, lists)
        e = ehp_of(seq, variants)
        itv = float(lists[lid]['enemyInterval'])
        print('W%02d  %5d  %9.0f  %8s  %8.2f' % (w, len(seq), e, '—' if prev is None else '%+.1f%%' % ((e / prev - 1) * 100), itv))
        prev = e


def compute_hash():
    keys = ['tbenemydata', 'tbenemylist', 'tbrounddata', 'tbsceneinfo',
            'tblevelmap', 'tbtowerinfo', 'tbbulletdata', 'tbglobal']
    text = []
    for k in keys:
        p = JSON_DIR / (k + '.json')
        s = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        text.append('%s\n%s\n' % (k, s))
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    rows.sort(key=lambda r: r['id'])
    for r in rows:
        def joined(v):
            return ','.join(str(x) for x in v) if isinstance(v, list) else ''
        fields = [r['id'], r['difficultyTier'], r['profileId'], r['profileRevision'],
                  r['expectedWaveCount'], r['expectedEnemyCount'], r['twoStarHpPermille'],
                  r['threeStarHpPermille'], r['fourStarHpPermille'], r['maxFourStarLeaks'],
                  r['fourStarTimePermille'], r['parCombatTimeMs'],
                  joined(r.get('bossVariantIds')), joined(r.get('criticalVariantIds')),
                  r['balanceRevision'], r['ratingRevision'], r['planHash']]
        text.append('|'.join(str(x) for x in fields) + '\n')
    return hashlib.sha256(''.join(text).encode('utf-8')).hexdigest()


def cmd_reseal():
    h = compute_hash()
    print('按导出 JSON 计算 runtimeContentHash = %s' % h)
    for i in range(1, 31):
        p = DATAS / ('Levels/L%02d/LevelRating.xlsx' % i)
        wb, ws, heads, rows = load_rows(p)
        for r, v in rows:
            ws.cell(r, heads['runtimeContentHash'], h)
        wb.save(p)
        wb.close()
    print('已写入 30 关 LevelRating.xlsx（runtimeContentHash）')


def cmd_verify():
    h = compute_hash()
    rows = json.loads((JSON_DIR / 'tblevelrating.json').read_text(encoding='utf-8-sig'))
    stored = sorted({r['runtimeContentHash'] for r in rows})
    print('重算哈希: %s' % h)
    print('表中哈希: %s 一致=%s' % (stored, stored == [h]))
    enemy = {x['id']: x for x in json.loads((JSON_DIR / 'tbenemydata.json').read_text(encoding='utf-8-sig'))}
    order, rounds, lists, _ = read_level()
    vals = {}
    for w in (19, 20, 21, 25):
        rid, lid, seq = seq_of(w, order, rounds, lists)
        vals[w] = sum(float(enemy[v]['hp']) / (1 - float(enemy[v]['armor'] or 0)) for v in seq)
    print('运行时 JSON：W19 %.0f → W20 %.0f（%+.2f%%）→ W21 %.0f（%+.2f%%）→ W25 %.0f'
          % (vals[19], vals[20], (vals[20] / vals[19] - 1) * 100,
             vals[21], (vals[21] / vals[20] - 1) * 100, vals[25]))
    itv = {w: float(lists[seq_of(w, order, rounds, lists)[1]]['enemyInterval']) for w in (19, 20, 21)}
    print('出怪间隔：W19 %.2f → W20 %.2f（%+.2f%%）→ W21 %.2f'
          % (itv[19], itv[20], (itv[20] / itv[19] - 1) * 100, itv[21]))


if __name__ == '__main__':
    mode = sys.argv[1] if len(sys.argv) > 1 else 'verify'
    {'values': cmd_values, 'reseal': cmd_reseal, 'verify': cmd_verify}[mode]()

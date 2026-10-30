# -*- coding: utf-8 -*-
"""L01 血量微调第四轮（2026-10-30 用户需求）：W15-25（含 15）地面怪 hp ×0.9（-10%），空中不动。

在 hp3-ground85 层之上叠加；W1-14、空中怪、数量/奖励/出怪/三塔均不动；
expectedEnemyCount/par 不变；LevelRating 仅 planHash/修订号更新。

用法：python l01_hp_tune4_w15.py [--write]
"""
import json, pathlib, sys
import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
K = 0.9                     # -10%
W_FROM, W_TO = 15, 25
TAG = 'hp4-20261030: W15-25 ground x0.9'
REVISION = 'hp4-ground90-20261030-v1'

def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    out = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            out.append({k: ws.cell(r, c).value for k, c in heads.items()})
    wb.close()
    return out

def load():
    scene = _rows(L01 / 'SceneInfo.xlsx')[0]
    rating = _rows(L01 / 'LevelRating.xlsx')[0]
    order = [int(x) for x in str(scene['RoundList']).split(',')]
    rounds = {int(r['id']): int(str(r['EnemyIndexs']).split(',')[0]) for r in _rows(L01 / 'RoundData.xlsx')}
    lists = {int(r['id']): dict(seq=[int(x) for x in str(r['EnemyIndexs']).split(',')],
                                desc=r['desc']) for r in _rows(L01 / 'EnemyList.xlsx')}
    variants = {int(r['id']): r for r in _rows(L01 / 'EnemyData.xlsx')}
    wmap = {rid: i + 1 for i, rid in enumerate(order)}
    for rid in order:
        w = wmap[rid]
        if w == 15:
            n = len(lists[rounds[rid]]['seq'])
            if n != 481:   # 460 + 另一会话 ground5-20261030 的地面 +5%（实测盘上值）
                raise RuntimeError(f'基线不符：W15 敌数 {n} != 481')
    if not any('hp3-20261030' in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('未检测到 hp3-20261030 层标记，拒绝执行')
    if any('hp4-20261030' in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('已应用本轮血量调整，拒绝重复执行')
    return dict(order=order, rounds=rounds, lists=lists, variants=variants, rating=rating)

def main(dry=True):
    lv = load()
    wmap = {rid: i + 1 for i, rid in enumerate(lv['order'])}
    buff_ids = set()
    for rid in lv['order']:
        w = wmap[rid]
        if W_FROM <= w <= W_TO:
            for v in lv['lists'][lv['rounds'][rid]]['seq']:
                if lv['variants'][v]['isFlying'] != 1:
                    buff_ids.add(v)
    new_hp = {v: round(float(lv['variants'][v]['hp']) * K, 6) for v in buff_ids}
    audit = []
    p = L01 / 'EnemyData.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    n = 0
    for r in range(4, ws.max_row + 1):
        vid = ws.cell(r, heads['id']).value
        if vid in buff_ids:
            ws.cell(r, heads['hp'], new_hp[vid]); n += 1
            d = str(ws.cell(r, heads['desc']).value or '')
            base = d.split(' | ' + TAG.split(':')[0])[0]
            ws.cell(r, heads['desc'], base + ' | ' + TAG)
    audit.append(('EnemyData.xlsx', f'{n} 个 W15-25 地面变体 hp ×0.9；空中/数量/奖励/出怪/W1-14 不动'))
    if not dry: wb.save(p)
    wb.close()
    p = L01 / 'EnemyList.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    for rid in lv['order']:
        w = wmap[rid]
        if not (W_FROM <= w <= W_TO): continue
        seq = lv['lists'][lv['rounds'][rid]]['seq']
        ehp = sum(new_hp.get(v, float(lv['variants'][v]['hp'])) / (1 - (lv['variants'][v]['armor'] or 0)) for v in seq)
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value == rid:
                d = str(ws.cell(r, heads['desc']).value or '')
                base = d.split(' | ' + TAG.split(':')[0])[0]
                ws.cell(r, heads['desc'], base + f' | {TAG} (EHP={ehp:.6f})')
                break
    audit.append(('EnemyList.xlsx', f'W{W_FROM}-W{W_TO} 共 11 行 desc 更新（序列/间隔不动）'))
    if not dry: wb.save(p)
    wb.close()
    import hashlib
    p = L01 / 'LevelRating.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    spec = json.dumps({'hp': new_hp}, ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value == 1:
            ws.cell(r, heads['planHash'], plan_hash)
            ws.cell(r, heads['balanceRevision'], REVISION)
    audit.append(('LevelRating.xlsx', f'planHash={plan_hash[:12]}…, revision={REVISION}（expectedEnemyCount/par 不变）'))
    if not dry: wb.save(p)
    wb.close()
    for a in audit:
        print(('（dry-run）' if dry else '已写入 '), a[0], ':', a[1])

if __name__ == '__main__':
    main(dry='--write' not in sys.argv)

# -*- coding: utf-8 -*-
"""L01 血量微调第二轮（2026-10-30 用户需求）：W5-25 地面 hp ×1.55、空中 hp ×1.05。

在 hp-20261030 层之上叠加，只改 W5-25 变体 hp；数量/奖励/出怪/W1-4/三塔不动；
expectedEnemyCount/par 不变；LevelRating 仅 planHash/修订号更新。

用法：python l01_hp_tune2.py [--write]
"""
import json, pathlib, sys, collections
import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
G, A = 1.55, 1.05
TAG = 'hp2-20261030: ground x1.55, air x1.05'
REVISION = 'hp2-55-5-20261030-v1'

def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    out = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            out.append({k: ws.cell(r, c).value for k, c in heads.items()})
    wb.close()
    return out

def load_level():
    scene = _rows(L01 / 'SceneInfo.xlsx')[0]
    rating = _rows(L01 / 'LevelRating.xlsx')[0]
    order = [int(x) for x in str(scene['RoundList']).split(',')]
    rounds = {int(r['id']): int(str(r['EnemyIndexs']).split(',')[0]) for r in _rows(L01 / 'RoundData.xlsx')}
    lists = {int(r['id']): dict(seq=[int(x) for x in str(r['EnemyIndexs']).split(',')],
                                desc=r['desc']) for r in _rows(L01 / 'EnemyList.xlsx')}
    variants = {int(r['id']): r for r in _rows(L01 / 'EnemyData.xlsx')}
    wmap = {rid: i + 1 for i, rid in enumerate(order)}
    for rid in order:
        w = wmap[rid]
        if w == 5:
            n = len(lists[rounds[rid]]['seq'])
            if n != 231:
                raise RuntimeError(f'基线不符：W5 敌数 {n} != 231')
    if not any('hp-20261030' in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('未检测到 hp-20261030 层标记，拒绝执行')
    if any('hp2-20261030' in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('已应用本轮血量调整，拒绝重复执行')
    return dict(order=order, rounds=rounds, lists=lists, variants=variants, rating=rating)

def compute(lv):
    wmap = {rid: i + 1 for i, rid in enumerate(lv['order'])}
    buff_ids = set()
    for rid in lv['order']:
        if wmap[rid] >= 5:
            buff_ids |= set(lv['lists'][lv['rounds'][rid]]['seq'])
    old_hp = {v: float(lv['variants'][v]['hp']) for v in buff_ids}
    new_hp = {v: round(hp * (A if lv['variants'][v]['isFlying'] == 1 else G), 6)
              for v, hp in old_hp.items()}
    waves = []
    for rid in lv['order']:
        w = wmap[rid]; seq = lv['lists'][lv['rounds'][rid]]['seq']
        ehp = sum((new_hp.get(v, lv['variants'][v]['hp']) / (1 - (lv['variants'][v]['armor'] or 0))) for v in seq)
        ehp_old = sum((old_hp.get(v, lv['variants'][v]['hp']) / (1 - (lv['variants'][v]['armor'] or 0))) for v in seq)
        waves.append(dict(wave=w, rid=rid, n=len(seq), ehp_old=ehp_old, ehp=ehp))
    return dict(buff_ids=buff_ids, old_hp=old_hp, new_hp=new_hp, waves=waves)

def write_tables(lv, res, dry=True):
    audit = []
    p = L01 / 'EnemyData.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    n_hp = 0
    for r in range(4, ws.max_row + 1):
        vid = ws.cell(r, heads['id']).value
        if vid in res['buff_ids']:
            ws.cell(r, heads['hp'], res['new_hp'][vid]); n_hp += 1
            d = str(ws.cell(r, heads['desc']).value or '')
            base = d.split(' | ' + TAG.split(':')[0])[0]
            mult = 'x1.05' if lv['variants'][vid]['isFlying'] == 1 else 'x1.55'
            ws.cell(r, heads['desc'], base + f' | {TAG} ({mult})')
    audit.append(('EnemyData.xlsx', f'{n_hp} 个 W5-25 变体 hp 更新（地面×1.55、空中×1.05）；数量/奖励/出怪/W1-4 不动'))
    if not dry: wb.save(p)
    wb.close()
    p = L01 / 'EnemyList.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    for x in res['waves']:
        if x['wave'] <= 4: continue
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value == x['rid']:
                d = str(ws.cell(r, heads['desc']).value or '')
                base = d.split(' | ' + TAG.split(':')[0])[0]
                ws.cell(r, heads['desc'], base + f' | {TAG} (EHP={x["ehp"]:.6f})')
                break
    audit.append(('EnemyList.xlsx', 'W5-25 共 21 行 desc 更新（序列/间隔不动）'))
    if not dry: wb.save(p)
    wb.close()
    import hashlib
    p = L01 / 'LevelRating.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    spec = json.dumps({'hp': res['new_hp']}, ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value == 1:
            ws.cell(r, heads['planHash'], plan_hash)
            ws.cell(r, heads['balanceRevision'], REVISION)
    audit.append(('LevelRating.xlsx', f'planHash={plan_hash[:12]}…, revision={REVISION}（expectedEnemyCount/par 不变）'))
    if not dry: wb.save(p)
    wb.close()
    return audit

if __name__ == '__main__':
    dry = '--write' not in sys.argv
    lv = load_level()
    res = compute(lv)
    print(f"地面×{G}  空中×{A}  变体数={len(res['buff_ids'])}")
    print('波 | EHP 旧→新')
    for x in res['waves']:
        if x['wave'] in (1, 4, 5, 10, 15, 20, 25):
            print(f"W{x['wave']:02d} | {x['ehp_old']:9.0f} -> {x['ehp']:9.0f}")
    for a in write_tables(lv, res, dry=dry):
        print(('（dry-run）' if dry else '已写入 '), a[0], ':', a[1])

# -*- coding: utf-8 -*-
"""L01 补丁（2026-10-30 用户追加）：W10 地面怪 hp 也 ×0.85（并入第三轮 −15% 范围）。

用法：python l01_hp_tune3b_w10.py [--write]
"""
import json, pathlib, sys, collections
import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
K = 0.85
TAG = 'hp3-20261030: W10 ground x0.85'
REVISION = 'hp3-ground85-20261030-v2'

def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    out = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            out.append({k: ws.cell(r, c).value for k, c in heads.items()})
    wb.close()
    return out

def load():
    scene = _rows(L01 / 'SceneInfo.xlsx')[0]
    rating = _rows(L01 / 'LevelRating.xlsx')[0]
    order = [int(x) for x in str(scene['RoundList']).split(',')]
    rounds = {int(r['id']): int(str(r['EnemyIndexs']).split(',')[0]) for r in _rows(L01 / 'RoundData.xlsx')}
    lists = {int(r['id']): dict(seq=[int(x) for x in str(r['EnemyIndexs']).split(',')],
                                desc=r['desc']) for r in _rows(L01 / 'EnemyList.xlsx')}
    variants = {int(r['id']): r for r in _rows(L01 / 'EnemyData.xlsx')}
    w10 = order[9]
    seq = lists[rounds[w10]]['seq']
    ground_ids = {v for v in set(seq) if variants[v]['isFlying'] != 1}
    already = any('hp3-20261030' in str(variants[v]['desc'] or '') for v in ground_ids)
    if already:
        raise RuntimeError('W10 地面变体已应用本轮调整，拒绝重复执行')
    return dict(order=order, rounds=rounds, lists=lists, variants=variants,
                w10_rid=w10, seq=seq, ground_ids=ground_ids, rating=rating)

def main(dry=True):
    lv = load()
    new_hp = {v: round(float(lv['variants'][v]['hp']) * K, 6) for v in lv['ground_ids']}
    old_ehp = sum(lv['variants'][v]['hp'] / (1 - (lv['variants'][v]['armor'] or 0)) for v in lv['seq'])
    new_ehp = sum(new_hp.get(v, lv['variants'][v]['hp']) / (1 - (lv['variants'][v]['armor'] or 0)) for v in lv['seq'])
    print(f'W10 地面变体 {len(lv["ground_ids"])} 个；波 EHP {old_ehp:.0f} -> {new_ehp:.0f}')
    audit = []
    p = L01 / 'EnemyData.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    n = 0
    for r in range(4, ws.max_row + 1):
        vid = ws.cell(r, heads['id']).value
        if vid in lv['ground_ids']:
            ws.cell(r, heads['hp'], new_hp[vid]); n += 1
            d = str(ws.cell(r, heads['desc']).value or '')
            base = d.split(' | ' + TAG.split(':')[0])[0]
            ws.cell(r, heads['desc'], base + ' | ' + TAG)
    audit.append(('EnemyData.xlsx', f'{n} 个 W10 地面变体 hp ×0.85'))
    if not dry: wb.save(p)
    wb.close()
    p = L01 / 'EnemyList.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value == lv['w10_rid']:
            d = str(ws.cell(r, heads['desc']).value or '')
            base = d.split(' | ' + TAG.split(':')[0])[0]
            ws.cell(r, heads['desc'], base + f' | {TAG} (EHP={new_ehp:.6f})')
            break
    audit.append(('EnemyList.xlsx', 'W10 desc 更新（序列/间隔不动）'))
    if not dry: wb.save(p)
    wb.close()
    import hashlib
    p = L01 / 'LevelRating.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    spec = json.dumps({'hp': new_hp}, ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value == 1:
            ws.cell(r, heads['planHash'], plan_hash)
            ws.cell(r, heads['balanceRevision'], REVISION)
    audit.append(('LevelRating.xlsx', f'planHash={plan_hash[:12]}…, revision={REVISION}'))
    if not dry: wb.save(p)
    wb.close()
    for a in audit:
        print(('（dry-run）' if dry else '已写入 '), a[0], ':', a[1])

if __name__ == '__main__':
    main(dry='--write' not in sys.argv)

# -*- coding: utf-8 -*-
"""L01 击杀金币奖励 +20%（2026-10-29 用户需求，在第二轮强化态上增量调整）：

- W5-25 击杀金预算 2040 → 2448（精确 +20%）；W1-4 与波次通关奖励不变。
- 分配：按当前奖励等比 ×1.2 取整，全局 ±1 局部搜索精确命中（移动步长 = 变体翻倍出场数）。
- 总收入 3259 → 3667。LevelRating 仅 planHash/修订号变化（敌数、par 不变）。
- 防重入护栏：基线必须是第二轮强化态（W5 N=220 且 desc 含 wavebuff2-20261029）。

用法：python l01_kill_gold_20.py [--write]
"""
import json, math, pathlib, sys, collections
import openpyxl

ROOT = pathlib.Path(r'D:/FishTestGame')
L01 = ROOT / 'Luban/Config/Datas/Levels/L01'
SCALE = 1.2
TAG = 'gold20-20261029'
REVISION = 'wavebuff2-gold20-20261029-v1'

def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    out = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value is not None:
            out.append({k: ws.cell(r, c).value for k, c in heads.items()})
    wb.close()
    return out

def load_level():
    scene = _rows(L01 / 'SceneInfo.xlsx')[0]
    rating = _rows(L01 / 'LevelRating.xlsx')[0]
    order = [int(x) for x in str(scene['RoundList']).split(',')]
    rounds = {int(r['id']): dict(list_id=int(str(r['EnemyIndexs']).split(',')[0]),
                                 reward=int(r['rewardGold'])) for r in _rows(L01 / 'RoundData.xlsx')}
    lists = {int(r['id']): dict(seq=[int(x) for x in str(r['EnemyIndexs']).split(',')],
                                enemyInterval=int(r['enemyInterval']), desc=r['desc'])
             for r in _rows(L01 / 'EnemyList.xlsx')}
    variants = {int(r['id']): r for r in _rows(L01 / 'EnemyData.xlsx')}
    wmap = {rid: i + 1 for i, rid in enumerate(order)}
    for rid in order:
        w = wmap[rid]
        n = len(lists[rounds[rid]['list_id']]['seq'])
        expect = 4 * (25 + 6 * w) if w >= 5 else 25 + 6 * w
        if n != expect:
            raise RuntimeError(f'基线不符：W{w} 敌数 {n} != {expect}（应为第二轮强化态）')
    if not any('wavebuff2-20261029' in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('未检测到第二轮强化标记，拒绝执行')
    if any(TAG in str(l['desc'] or '') for l in lists.values()):
        raise RuntimeError('已应用本次 +20% 调整，拒绝重复执行')
    return dict(order=order, rounds=rounds, lists=lists, variants=variants,
                initialGold=int(scene['initialGold']), rating=rating)

def compute(lv):
    wmap = {rid: i + 1 for i, rid in enumerate(lv['order'])}
    cnt4 = collections.Counter()          # 当前序列中的出场次数（本轮数量不变，仅奖励 ×1.2）
    old_kill_w = {}
    for rid in lv['order']:
        w = wmap[rid]; seq = lv['lists'][lv['rounds'][rid]['list_id']]['seq']
        kill = sum(int(lv['variants'][v]['reward'] or 0) for v in seq)
        old_kill_w[w] = kill
        if w >= 5:
            for v in seq: cnt4[v] += 1
    kill_w14 = sum(old_kill_w[w] for w in range(1, 5))
    old_w525 = sum(old_kill_w[w] for w in range(5, 26))
    target_w525 = round(old_w525 * SCALE / 4) * 4      # 移动步长为 4 的倍数，目标也取 4 的倍数
    new_reward = {v: min(9, max(0, round(int(lv['variants'][v]['reward'] or 0) * SCALE))) for v in cnt4}
    def kill_sum(): return sum(new_reward[v] * cnt4[v] for v in cnt4)
    diff = target_w525 - kill_sum()
    guard = 0
    while diff != 0 and guard < 500:
        guard += 1
        best = None
        for v in cnt4:
            if new_reward[v] < 9:
                nd = abs(diff - cnt4[v])
                if best is None or nd < best[0]: best = (nd, v, +1)
            if new_reward[v] > 0:
                nd = abs(diff + cnt4[v])
                if nd < best[0]: best = (nd, v, -1)
        if best is None or best[0] >= abs(diff): break
        _, v, step = best
        new_reward[v] += step; diff -= step * cnt4[v]
    assert abs(diff) <= 10, f'精确和残差过大：{diff}'
    waves = []
    for rid in lv['order']:
        w = wmap[rid]; seq = lv['lists'][lv['rounds'][rid]['list_id']]['seq']
        kill = sum(new_reward.get(v, int(lv['variants'][v]['reward'] or 0)) for v in seq)
        waves.append(dict(wave=w, rid=rid, n=len(seq), kill=kill, reward=lv['rounds'][rid]['reward']))
    income_new = lv['initialGold'] + sum(x['kill'] for x in waves) + sum(x['reward'] for x in waves)
    return dict(waves=waves, new_reward=new_reward, cnt4=cnt4,
                kill_w14=kill_w14, old_w525=old_w525, target_w525=target_w525,
                income_old=lv['initialGold'] + sum(old_kill_w.values()) + sum(x['reward'] for x in waves),
                income_new=income_new, diff=diff)

def write_tables(lv, res, dry=True):
    audit = []
    p = L01 / 'EnemyData.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    n_rw = 0
    for r in range(4, ws.max_row + 1):
        vid = ws.cell(r, heads['id']).value
        if vid in res['new_reward'] and ws.cell(r, heads['reward']).value != res['new_reward'][vid]:
            ws.cell(r, heads['reward'], res['new_reward'][vid]); n_rw += 1
            d = str(ws.cell(r, heads['desc']).value or '')
            base = d.split(' | ' + TAG)[0]
            ws.cell(r, heads['desc'], base + f' | {TAG}: reward +20%')
    audit.append(('EnemyData.xlsx', f'{n_rw} 个变体 reward 调整（×1.2 等比+精确和），hp/数量不动'))
    if not dry: wb.save(p)
    wb.close()
    p = L01 / 'EnemyList.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    for x in res['waves']:
        if x['wave'] <= 4: continue
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, heads['id']).value == x['rid']:
                d = str(ws.cell(r, heads['desc']).value or '')
                base = d.split(' | ' + TAG)[0]
                ws.cell(r, heads['desc'], base + f' | {TAG}: killGold={x["kill"]}')
                break
    audit.append(('EnemyList.xlsx', 'W5-25 desc 追加新击杀金（序列/间隔不动）'))
    if not dry: wb.save(p)
    wb.close()
    import hashlib
    p = L01 / 'LevelRating.xlsx'
    wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
    heads = {c.value: c.column for c in ws[1] if c.value}
    spec = json.dumps({'waves': [{k: x[k] for k in ('wave', 'kill')} for x in res['waves']],
                       'reward': res['new_reward']}, ensure_ascii=False, sort_keys=True)
    plan_hash = hashlib.sha256(spec.encode('utf-8')).hexdigest()
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, heads['id']).value == 1:
            ws.cell(r, heads['planHash'], plan_hash)
            ws.cell(r, heads['balanceRevision'], REVISION)
    audit.append(('LevelRating.xlsx', f'planHash={plan_hash[:12]}…, revision={REVISION}（敌数/par 不变）'))
    if not dry: wb.save(p)
    wb.close()
    return audit

if __name__ == '__main__':
    dry = '--write' not in sys.argv
    lv = load_level()
    res = compute(lv)
    print(f"W5-25 击杀金：{res['old_w525']} -> {res['target_w525']}（+{(res['target_w525']/res['old_w525']-1)*100:.1f}%）  "
          f"W1-4 击杀金 {res['kill_w14']} 不变")
    print(f"总收入：{res['income_old']} -> {res['income_new']}（{res['income_new']/res['income_old']-1:+.1%}）  残差 {res['diff']}")
    for x in res['waves']:
        if x['wave'] in (5, 10, 15, 20, 24, 25):
            print(f"  W{x['wave']:02d} 击杀金 -> {x['kill']}")
    for a in write_tables(lv, res, dry=dry):
        print(('（dry-run）' if dry else '已写入 '), a[0], ':', a[1])
